"""Fixes for what real conversions turned up (2026-10-24), fully offline.

* A flowchart decision diamond had no native kind, so the model cropped it out
  of the source as an ``image`` and then drew its label again on top.
* ``.env.example``'s placeholder key reached Google, whose "invalid key" 400 was
  retried without JSON mode and reported as a malformed request.
"""

import importlib
from types import SimpleNamespace

import pytest
from pptx.enum.shapes import MSO_SHAPE

from app import config, vision_analyzer
from app.pptx_builder import build_presentation
from app.schemas import SlideSpec
from app.services.preview import build_preview


class TestDiamond:
    @pytest.mark.parametrize("spelling", ["diamond", "rhombus", "decision"])
    def test_every_spelling_builds_a_native_diamond(self, spelling):
        spec = SlideSpec.model_validate(
            {"elements": [{"type": "shape", "kind": spelling, "bbox": [300, 600, 160, 180]}]}
        )
        assert spec.elements[0].kind == "diamond"
        shape = build_presentation(spec).slides[0].shapes[0]
        assert shape.auto_shape_type == MSO_SHAPE.DIAMOND

    def test_the_preview_names_it(self):
        spec = SlideSpec.model_validate(
            {"elements": [{"type": "shape", "kind": "diamond", "bbox": [300, 600, 160, 180]}]}
        )
        preview = build_preview(build_presentation(spec))
        shapes = [e for e in preview["shapes"] if e.get("kind") == "shape"]
        assert [s["autoshape"] for s in shapes] == ["diamond"]

    def test_the_prompt_offers_it_and_forbids_cropping_shapes(self):
        prompt = vision_analyzer.SYSTEM_PROMPT
        assert '"diamond"' in prompt
        assert "ALWAYS a \"shape\", never an \"image\"" in prompt


class TestPlaceholderKeys:
    _VARS = ("PROVIDER", "LITELLM_API_KEY", "GEMINI_API_KEY", "OPENAI_API_KEY",
             "ANTHROPIC_API_KEY")

    @pytest.fixture
    def reload_with(self, monkeypatch):
        def _reload(**env):
            for name in self._VARS:
                monkeypatch.delenv(name, raising=False)
            for name, value in env.items():
                monkeypatch.setenv(name, value)
            return importlib.reload(config)

        yield _reload
        for name in self._VARS:
            monkeypatch.delenv(name, raising=False)
        importlib.reload(config)

    @pytest.mark.parametrize(
        "placeholder",
        ["replace-with-your-gemini-key", "<your gemini key>", "sk-replace-with-your-own-key"],
    )
    def test_an_unedited_placeholder_is_a_missing_key(self, reload_with, placeholder):
        reloaded = reload_with(PROVIDER="gemini", GEMINI_API_KEY=placeholder)
        assert reloaded.LITELLM_API_KEY == ""
        with pytest.raises(reloaded.MissingApiKeyError):
            reloaded.get_client()

    def test_a_placeholder_does_not_win_the_pick_by_key(self, reload_with):
        reloaded = reload_with(
            GEMINI_API_KEY="replace-with-your-gemini-key", OPENAI_API_KEY="o-real"
        )
        assert (reloaded.PROVIDER, reloaded.LITELLM_API_KEY) == ("openai", "o-real")


class TestABadKeyIsNotRetried:
    def test_googles_invalid_key_400_is_reported_as_a_key_problem(self):
        calls = []

        def create(**kwargs):
            calls.append(kwargs)
            raise RuntimeError(
                "BadRequestError: Error code: 400 - [{'error': {'code': 400, "
                "'message': 'Please pass a valid API key', 'status': 'INVALID_ARGUMENT'}}]"
            )

        client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
        with pytest.raises(vision_analyzer.VisionAnalysisError) as excinfo:
            vision_analyzer._create_completion(client, "gemini-3.8-flash", [])
        assert len(calls) == 1
        assert "rejected the API key" in str(excinfo.value)
        assert "response_format" not in str(excinfo.value)
