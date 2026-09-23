"""CLI tests plus the sample-image generator - all offline.

``main()`` is called in-process and its *return code* is asserted, because the
exit-code contract (0 ok / 1 runtime failure / 2 bad usage) is the part other
tooling depends on.  ``_run`` also catches ``SystemExit`` so an argparse-driven
failure path is graded the same way as a returned code.
"""

from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from app import config
import cli
from app import mock_data
from app import vision_analyzer
from make_sample_image import make_sample_image
from app.schemas import ShapeElement, SlideSpec, TableElement, TextElement
from app.vision_analyzer import VisionAnalysisError


def _run(argv: list[str]) -> int:
    """Call ``main(argv)`` and normalize argparse's ``SystemExit`` to an int."""
    try:
        return cli.main(argv)
    except SystemExit as exc:  # argparse usage errors
        code = exc.code
        return 0 if code is None else int(code)


def _patch_analyzer(monkeypatch, fake):
    """Replace ``analyze_image`` wherever main might have bound it."""
    monkeypatch.setattr(vision_analyzer, "analyze_image", fake, raising=False)
    if hasattr(cli, "analyze_image"):
        monkeypatch.setattr(cli, "analyze_image", fake)


def _forbid_network(monkeypatch):
    """Make any attempt to build a real LLM client an outright failure."""

    def boom(*args, **kwargs):
        raise AssertionError("the LLM path must not be touched here")

    monkeypatch.setattr(config, "get_client", boom, raising=False)
    if hasattr(cli, "get_client"):
        monkeypatch.setattr(cli, "get_client", boom)
    _patch_analyzer(monkeypatch, boom)


def _sample_png(tmp_path, name: str = "sample.png") -> Path:
    return make_sample_image(tmp_path / name, size=(320, 180))


class _ReplyClient:
    """Minimal ``openai.OpenAI`` stand-in that always answers with one reply.

    Lets a test drive the REAL analyzer (and therefore the real parse path) off a
    chosen model reply, without touching the network.
    """

    def __init__(self, reply: str) -> None:
        self.reply = reply
        client = self

        class _Completions:
            def create(self, **kwargs):
                message = type("_Msg", (), {"content": client.reply})()
                choice = type("_Choice", (), {"message": message, "finish_reason": "stop"})()
                return type("_Response", (), {"choices": [choice]})()

        self.chat = type("_Chat", (), {"completions": _Completions()})()


def _assert_real_deck(path: Path) -> Presentation:
    prs = Presentation(str(path))
    assert len(prs.slides) == 1
    assert len(prs.slides[0].shapes) > 0
    return prs


# ---------------------------------------------------------------------------
# validate_input_path
# ---------------------------------------------------------------------------
class TestValidateInputPath:
    def test_accepts_a_supported_image(self, tmp_path):
        png = _sample_png(tmp_path)
        result = cli.validate_input_path(str(png))
        assert isinstance(result, Path)
        assert result.exists()
        assert result.name == png.name

    @pytest.mark.parametrize("name", ["notes.bmp", "scan.tiff", "spec.pdf", "data.txt"])
    def test_rejects_unsupported_extensions(self, tmp_path, name):
        path = tmp_path / name
        path.write_bytes(b"not really an image")
        with pytest.raises((Exception, SystemExit)):
            cli.validate_input_path(str(path))

    @pytest.mark.parametrize("name", [".png", ".jpg", ".webp"])
    def test_a_dotfile_name_is_not_quoted_back_as_its_own_extension(
        self, tmp_path, name
    ):
        """``Path(".png").suffix`` is '' - POSIX reads a leading-dot name as a
        name, not a suffix - so quoting the filename produced "Unsupported image
        type '.png'. Supported extensions: ..., .png, ...", which rejects and
        allows the same thing. The message has to describe what is actually
        wrong: the name carries no extension."""
        path = tmp_path / name
        path.write_bytes(b"not really an image")
        with pytest.raises(cli.CliUsageError) as caught:
            cli.validate_input_path(str(path))
        message = str(caught.value)
        rejected = message.split("Supported extensions:")[0]
        assert name not in rejected
        assert "<no extension>" in rejected
        for extension in sorted(config.SUPPORTED_EXTENSIONS):
            assert extension in message
        assert "slide.png" in message

    def test_a_real_unsupported_extension_is_still_named(self, tmp_path):
        path = tmp_path / "scan.tiff"
        path.write_bytes(b"not really an image")
        with pytest.raises(cli.CliUsageError) as caught:
            cli.validate_input_path(str(path))
        assert "'.tiff'" in str(caught.value)

    def test_rejects_a_missing_file(self, tmp_path):
        with pytest.raises((Exception, SystemExit)):
            cli.validate_input_path(str(tmp_path / "ghost.png"))

    def test_rejects_a_directory(self, tmp_path):
        directory = tmp_path / "shots.png"
        directory.mkdir()
        with pytest.raises((Exception, SystemExit)):
            cli.validate_input_path(str(directory))

    def test_rejects_a_non_image_wearing_an_image_extension(self, tmp_path):
        """A renamed file / HTML error page must fail locally, not at the proxy."""
        fake = tmp_path / "not-really.png"
        fake.write_bytes(b"<html>404 Not Found</html>")
        with pytest.raises(cli.CliUsageError):
            cli.validate_input_path(str(fake))

    def test_rejects_a_truncated_image(self, tmp_path):
        """A half-downloaded PNG is corrupt; catching it early saves two uploads."""
        good = _sample_png(tmp_path, "good.png")
        truncated = tmp_path / "half.png"
        truncated.write_bytes(good.read_bytes()[:300])
        with pytest.raises(cli.CliUsageError):
            cli.validate_input_path(str(truncated))

    def test_a_corrupt_image_exits_two_without_touching_the_llm(
        self, tmp_path, monkeypatch
    ):
        _forbid_network(monkeypatch)
        fake = tmp_path / "broken.png"
        fake.write_bytes(b"definitely not a PNG")
        assert _run(["-i", str(fake), "-o", str(tmp_path / "o.pptx")]) == 2


# ---------------------------------------------------------------------------
# --mock is gone: every case it used to cover, now on the --from-json path
# ---------------------------------------------------------------------------
def _spec_file(tmp_path, name: str = "spec.json") -> Path:
    """The sample spec on disk - the offline, no-key source ``--mock`` used to be."""
    path = tmp_path / name
    path.write_text(mock_data.SAMPLE_SPEC_JSON, encoding="utf-8")
    return path


class TestMockFlagIsGone:
    """``--mock`` was a documented flag; it is not one any more.

    argparse refuses an unknown option by itself, which is exactly the desired
    behaviour (exit 2, a usage error naming the argument) - but only as long as
    nobody adds the flag back as a silent no-op, so it is asserted.
    """

    def test_the_flag_exits_two_with_an_unrecognized_argument_error(
        self, tmp_path, capsys
    ):
        assert _run(["--mock", "-o", str(tmp_path / "x.pptx")]) == 2
        err = capsys.readouterr().err
        assert "unrecognized arguments" in err
        assert "--mock" in err
        assert not (tmp_path / "x.pptx").exists()

    def test_the_flag_is_not_in_the_parser_at_all(self):
        parser = cli.build_parser()
        options = {
            option for action in parser._actions for option in action.option_strings
        }
        assert "--mock" not in options
        # The two sources that remain are still there.
        assert {"-i", "--input", "--from-json"} <= options
        assert not hasattr(parser.parse_args([]), "mock")

    def test_the_help_text_never_mentions_mock(self, capsys):
        assert _run(["--help"]) == 0
        printed = capsys.readouterr().out
        assert "mock" not in printed.lower()
        assert "--from-json" in printed

    def test_the_no_input_usage_error_never_mentions_mock(self, capsys):
        assert _run([]) == 2
        message = "".join(capsys.readouterr()).lower()
        assert "mock" not in message
        assert "--input" in message and "--from-json" in message


class TestTheOfflineBuildPath:
    """``--from-json`` now carries what ``--mock`` used to: a full build with no
    network, no key and no model. Each case below is the one ``TestMockMode``
    asserted, moved onto the flag that still exists."""

    def test_returns_zero_and_writes_a_deck(self, tmp_path):
        out = tmp_path / "offline.pptx"
        assert _run(["--from-json", str(_spec_file(tmp_path)), "-o", str(out)]) == 0
        assert out.exists() and out.stat().st_size > 0
        _assert_real_deck(out)

    def test_never_touches_the_network(self, tmp_path, monkeypatch):
        _forbid_network(monkeypatch)
        out = tmp_path / "offline.pptx"
        assert _run(["--from-json", str(_spec_file(tmp_path)), "-o", str(out)]) == 0
        assert out.exists()

    def test_deck_matches_the_sample_spec_element_count(self, tmp_path):
        out = tmp_path / "offline.pptx"
        assert _run(["--from-json", str(_spec_file(tmp_path)), "-o", str(out)]) == 0
        spec = mock_data.sample_spec()
        # >= rather than ==: a builder is free to realise the slide background
        # as one extra shape.  Exact per-element mapping is asserted in
        # tests/test_pptx_builder.py against a spec with no background colour.
        assert len(Presentation(str(out)).slides[0].shapes) >= len(spec.elements)

    def test_prints_a_summary_with_the_output_path_and_counts(self, tmp_path, capsys):
        out = tmp_path / "offline.pptx"
        assert _run(["--from-json", str(_spec_file(tmp_path)), "-o", str(out)]) == 0
        printed = "".join(capsys.readouterr()).lower()
        assert out.name.lower() in printed
        assert any(word in printed for word in ("shape", "text", "table"))

    def test_creates_missing_output_directories(self, tmp_path):
        out = tmp_path / "deep" / "deeper" / "offline.pptx"
        assert _run(["--from-json", str(_spec_file(tmp_path)), "-o", str(out)]) == 0
        assert out.exists()


# ---------------------------------------------------------------------------
# --from-json
# ---------------------------------------------------------------------------
class TestFromJson:
    def test_builds_a_deck_from_a_saved_spec_without_an_llm(self, tmp_path, monkeypatch):
        _forbid_network(monkeypatch)
        spec_path = tmp_path / "spec.json"
        spec_path.write_text(mock_data.SAMPLE_SPEC_JSON, encoding="utf-8")
        out = tmp_path / "from_json.pptx"

        assert _run(["--from-json", str(spec_path), "-o", str(out)]) == 0
        prs = _assert_real_deck(out)
        assert len(prs.slides[0].shapes) == len(mock_data.sample_spec().elements)

    def test_invalid_json_file_fails_with_a_nonzero_code(self, tmp_path, capsys):
        spec_path = tmp_path / "broken.json"
        spec_path.write_text("{not json at all", encoding="utf-8")
        code = _run(["--from-json", str(spec_path), "-o", str(tmp_path / "x.pptx")])
        assert code in (1, 2)
        assert capsys.readouterr().err.strip(), "a failure must explain itself on stderr"

    def test_missing_json_file_fails_with_a_nonzero_code(self, tmp_path):
        assert _run(["--from-json", str(tmp_path / "ghost.json"), "-o", str(tmp_path / "x.pptx")]) in (1, 2)

    @pytest.mark.parametrize(
        "payload",
        [
            '{"slide": {"elements": [{"type": "text", "bbox": [0, 0, 500, 100], '
            '"content": "Recovered"}]}}',
            '[{"type": "text", "bbox": [0, 0, 500, 100], "content": "Recovered"}]',
        ],
    )
    def test_an_envelope_or_bare_array_spec_is_recovered(self, tmp_path, payload):
        """Neither shape may validate into a silently empty (blank) deck."""
        spec_path = tmp_path / "spec.json"
        spec_path.write_text(payload, encoding="utf-8")
        out = tmp_path / "recovered.pptx"

        assert _run(["--from-json", str(spec_path), "-o", str(out)]) == 0
        slide = _assert_real_deck(out).slides[0]
        assert any(
            "Recovered" in s.text_frame.text for s in slide.shapes if s.has_text_frame
        )

    def test_an_empty_spec_is_reported_loudly(self, tmp_path, capsys):
        spec_path = tmp_path / "empty.json"
        spec_path.write_text('{"background_color": "#FFFFFF", "elements": []}', encoding="utf-8")
        out = tmp_path / "empty.pptx"

        code = _run(["--from-json", str(spec_path), "-o", str(out)])
        captured = capsys.readouterr()
        assert "blank" in captured.err.lower(), "a blank deck must be called out on stderr"
        assert "fully editable" not in captured.out, "do not claim success for a blank slide"
        assert code == 0 and out.exists()


# ---------------------------------------------------------------------------
# the analyze path, with the LLM replaced
# ---------------------------------------------------------------------------
class TestAnalyzePath:
    def _fake(self, recorder):
        def fake_analyze(
            image_path, model=None, client=None, extra_instructions=None, verbose=False
        ):
            recorder["image_path"] = Path(str(image_path))
            recorder["model"] = model
            recorder["extra"] = extra_instructions
            recorder["verbose"] = verbose
            return mock_data.sample_spec()

        return fake_analyze

    def test_end_to_end_with_a_stubbed_analyzer(self, tmp_path, monkeypatch):
        recorder: dict = {}
        _patch_analyzer(monkeypatch, self._fake(recorder))
        png = _sample_png(tmp_path)
        out = tmp_path / "result.pptx"

        assert _run(["--input", str(png), "--output", str(out)]) == 0
        assert recorder["image_path"].name == png.name
        _assert_real_deck(out)

    def test_model_and_extra_flags_are_forwarded(self, tmp_path, monkeypatch):
        recorder: dict = {}
        _patch_analyzer(monkeypatch, self._fake(recorder))
        png = _sample_png(tmp_path)

        code = _run(
            [
                "-i",
                str(png),
                "-o",
                str(tmp_path / "r.pptx"),
                "-m",
                "my-vision-model",
                "--extra",
                "the header band is navy",
            ]
        )
        assert code == 0
        assert recorder["model"] == "my-vision-model"
        assert recorder["extra"] == "the header band is navy"

    def test_output_path_defaults_to_the_input_stem(self, tmp_path, monkeypatch):
        recorder: dict = {}
        _patch_analyzer(monkeypatch, self._fake(recorder))
        monkeypatch.chdir(tmp_path)
        png = _sample_png(tmp_path, "capture.png")

        assert _run(["-i", str(png)]) == 0
        assert (tmp_path / "capture.pptx").exists()

    def test_save_json_dumps_a_reusable_spec(self, tmp_path, monkeypatch):
        recorder: dict = {}
        _patch_analyzer(monkeypatch, self._fake(recorder))
        png = _sample_png(tmp_path)
        out = tmp_path / "r.pptx"
        dump = tmp_path / "spec.json"

        assert _run(["-i", str(png), "-o", str(out), "--save-json", str(dump)]) == 0
        assert dump.exists(), "--save-json must write the validated spec"

        payload = json.loads(dump.read_text(encoding="utf-8"))
        reparsed = SlideSpec.model_validate(payload)
        assert len(reparsed.elements) == len(mock_data.sample_spec().elements)
        # and the dump must be good enough to rebuild from
        assert _run(["--from-json", str(dump), "-o", str(tmp_path / "again.pptx")]) == 0

    def test_llm_failure_returns_one_with_a_friendly_message(self, tmp_path, monkeypatch, capsys):
        def boom(*args, **kwargs):
            raise VisionAnalysisError(
                "could not reach the LiteLLM proxy at http://localhost:4000/"
            )

        _patch_analyzer(monkeypatch, boom)
        png = _sample_png(tmp_path)

        assert _run(["-i", str(png), "-o", str(tmp_path / "r.pptx")]) == 1
        err = capsys.readouterr().err
        assert err.strip(), "the failure must be explained on stderr"
        assert "Traceback" not in err, "tracebacks are for --verbose only"

    def test_verbose_is_forwarded_to_the_analyzer(self, tmp_path, monkeypatch):
        """-v promises payload diagnostics, which only the analyzer can print."""
        recorder: dict = {}
        _patch_analyzer(monkeypatch, self._fake(recorder))
        png = _sample_png(tmp_path)

        assert _run(["-i", str(png), "-o", str(tmp_path / "r.pptx"), "-v"]) == 0
        assert recorder["verbose"] is True

    def test_verbose_defaults_to_off(self, tmp_path, monkeypatch):
        recorder: dict = {}
        _patch_analyzer(monkeypatch, self._fake(recorder))
        png = _sample_png(tmp_path)

        assert _run(["-i", str(png), "-o", str(tmp_path / "r.pptx")]) == 0
        assert recorder["verbose"] is False

    def test_verbose_prints_the_request_payload_and_the_raw_reply(
        self, tmp_path, monkeypatch, capsys
    ):
        """End to end through the real analyzer with a fake client."""

        class _Msg:
            content = mock_data.SAMPLE_SPEC_JSON

        class _Choice:
            message = _Msg()
            finish_reason = "stop"

        class _Response:
            choices = [_Choice()]

        class _Completions:
            def create(self, **kwargs):
                return _Response()

        class _Client:
            chat = type("_Chat", (), {"completions": _Completions()})()

        monkeypatch.setattr(config, "get_client", lambda: _Client())
        png = _sample_png(tmp_path)

        assert _run(["-i", str(png), "-o", str(tmp_path / "r.pptx"), "-v"]) == 0
        err = capsys.readouterr().err
        assert "request payload" in err, "-v must show what was sent"
        assert "raw model reply" in err, "-v must show what came back"
        assert "base64 chars omitted" in err, "-v must not dump the image bytes"

    def test_a_reply_without_elements_fails_instead_of_writing_a_blank_deck(
        self, tmp_path, monkeypatch, capsys
    ):
        def fake_analyze(image_path, model=None, client=None, extra_instructions=None, verbose=False):
            raise VisionAnalysisError("returned no usable slide elements")

        _patch_analyzer(monkeypatch, fake_analyze)
        png = _sample_png(tmp_path)
        out = tmp_path / "blank.pptx"

        assert _run(["-i", str(png), "-o", str(out)]) == 1
        assert not out.exists(), "a blank deck must not be written"
        assert capsys.readouterr().err.strip()

    def test_verbose_shows_the_traceback(self, tmp_path, monkeypatch, capsys):
        def boom(*args, **kwargs):
            raise VisionAnalysisError("kaput")

        _patch_analyzer(monkeypatch, boom)
        png = _sample_png(tmp_path)

        assert _run(["-i", str(png), "-o", str(tmp_path / "r.pptx"), "-v"]) == 1
        assert capsys.readouterr().err.strip()

    def test_multiline_error_details_stay_indented_under_the_headline(
        self, tmp_path, monkeypatch, capsys
    ):
        """A multi-line exception message must not spill back to column 0."""

        def boom(*args, **kwargs):
            raise VisionAnalysisError(
                "Vision request failed.\n"
                "  first attempt: boom\n"
                "Check that the proxy is running."
            )

        _patch_analyzer(monkeypatch, boom)
        png = _sample_png(tmp_path)
        assert _run(["-i", str(png), "-o", str(tmp_path / "r.pptx")]) == 1

        lines = [ln for ln in capsys.readouterr().err.splitlines() if ln.strip()]
        headline = next(i for i, ln in enumerate(lines) if ln.startswith(f"{cli.PROG}: error:"))
        detail = lines[headline + 1 :]
        assert detail, "the failure must carry detail lines"
        assert all(ln.startswith("  ") for ln in detail), detail


# ---------------------------------------------------------------------------
# partial parse: some elements unreadable, the deck still gets built
# ---------------------------------------------------------------------------
_GOOD_TEXT = {"type": "text", "bbox": [1, 1, 400, 80], "content": "good one"}
_GOOD_SHAPE = {
    "type": "shape",
    "bbox": [2, 200, 300, 120],
    "kind": "circle",
    "fill_color": "#FF0000",
}
#: A table element the way a vision model gets it wrong: no "data" key.
_TABLE_NO_DATA = {"type": "table", "bbox": [1, 400, 900, 400], "header_bg_color": "#112233"}


def _partial_spec() -> SlideSpec:
    spec = SlideSpec.model_validate(
        {"elements": [_GOOD_TEXT, _TABLE_NO_DATA, _GOOD_SHAPE]}
    )
    assert len(spec.elements) == 2 and len(spec.dropped) == 1, "precondition"
    return spec


class TestPartialDropIsReportedButNotFatal:
    def test_the_cli_warns_on_stderr_and_still_exits_zero(self, tmp_path, monkeypatch, capsys):
        spec = _partial_spec()
        _patch_analyzer(
            monkeypatch,
            lambda image_path, model=None, client=None, extra_instructions=None, verbose=False: spec,
        )
        png = _sample_png(tmp_path)
        out = tmp_path / "partial.pptx"

        assert _run(["-i", str(png), "-o", str(out)]) == 0, "a usable deck is a success"
        captured = capsys.readouterr()

        assert "1 element skipped" in captured.err
        assert "element 1 (table): data - Field required" in captured.err
        assert "skipped" not in captured.out, "stdout must stay clean"
        assert "Wrote" in captured.out
        prs = _assert_real_deck(out)
        assert len(prs.slides[0].shapes) == 2, "both good elements must be drawn"

    def test_the_saved_json_is_a_clean_spec_that_from_json_re_reads(
        self, tmp_path, monkeypatch, capsys
    ):
        spec = _partial_spec()
        _patch_analyzer(
            monkeypatch,
            lambda image_path, model=None, client=None, extra_instructions=None, verbose=False: spec,
        )
        png = _sample_png(tmp_path)
        dump = tmp_path / "spec.json"

        assert _run(["-i", str(png), "-o", str(tmp_path / "a.pptx"), "--save-json", str(dump)]) == 0
        capsys.readouterr()

        payload = json.loads(dump.read_text(encoding="utf-8"))
        assert "dropped" not in payload, "diagnostics must not pollute the spec file"
        assert len(payload["elements"]) == 2

        assert _run(["--from-json", str(dump), "-o", str(tmp_path / "b.pptx")]) == 0
        rebuilt = capsys.readouterr()
        assert "skipped" not in rebuilt.err, "a clean spec re-reads with nothing dropped"
        assert len(_assert_real_deck(tmp_path / "b.pptx").slides[0].shapes) == 2

    def test_a_spec_file_with_one_broken_element_still_builds(self, tmp_path, capsys):
        spec_path = tmp_path / "partial.json"
        spec_path.write_text(
            json.dumps({"elements": [_GOOD_TEXT, _TABLE_NO_DATA, _GOOD_SHAPE]}),
            encoding="utf-8",
        )
        out = tmp_path / "from_partial.pptx"

        assert _run(["--from-json", str(spec_path), "-o", str(out)]) == 0
        err = capsys.readouterr().err
        assert "1 element skipped" in err
        assert str(spec_path) in err, "the warning must name the source it read"
        assert len(_assert_real_deck(out).slides[0].shapes) == 2

    def test_a_spec_file_whose_every_element_is_broken_fails_loudly(self, tmp_path, capsys):
        spec_path = tmp_path / "hopeless.json"
        spec_path.write_text(
            json.dumps({"elements": [_TABLE_NO_DATA, {"type": "hologram", "bbox": [0, 0, 1, 1]}]}),
            encoding="utf-8",
        )
        out = tmp_path / "never.pptx"

        assert _run(["--from-json", str(spec_path), "-o", str(out)]) == 1
        err = capsys.readouterr().err
        assert "element 0 (table): data - Field required" in err
        assert not out.exists(), "a blank deck must not be written"


# ---------------------------------------------------------------------------
# the four end-to-end reproductions, as the user hit them
# ---------------------------------------------------------------------------
def _from_json(tmp_path, payload, name: str = "spec.json"):
    """Write *payload* (str or JSON-able) and return ``(exit_code, out_path)``."""
    spec_path = tmp_path / name
    text = payload if isinstance(payload, str) else json.dumps(payload)
    spec_path.write_text(text, encoding="utf-8")
    out = tmp_path / f"{spec_path.stem}.pptx"
    return _run(["--from-json", str(spec_path), "-o", str(out)]), out


class TestAnUnreadableBboxIsNeverSilent:
    """The CLI reported success, with an EMPTY stderr, over an invisible sliver.

    ``bbox: {"x":100,"y":200,"w":600}`` defaulted ``h`` to 1 and drew a text box
    8.000 x 0.050 in - 17 characters of 12pt text in 0.05in of height - while
    stdout said "2 elements ... fully editable".
    """

    #: Every bbox shape from the report, all of them silently accepted before.
    SLIVERS = [
        {"x": 100, "y": 200, "w": 600},
        [],
        [5, 5],
        [100, 200, 600],
        {"x": 100, "y": 200},
        {"q": 1},
        {"left_pct": 10},
        {"x": None, "y": None, "w": None, "h": None},
    ]

    @pytest.mark.parametrize("bad_bbox", SLIVERS, ids=lambda b: json.dumps(b))
    def test_the_element_is_dropped_named_and_still_exits_zero(
        self, tmp_path, capsys, bad_bbox
    ):
        code, out = _from_json(
            tmp_path,
            {
                "elements": [
                    {"type": "text", "bbox": bad_bbox, "content": "QUARTERLY REVENUE"},
                    _GOOD_SHAPE,
                ]
            },
        )
        captured = capsys.readouterr()

        assert code == 0, "the good shape still builds, so this is a partial success"
        assert "1 element skipped" in captured.err, captured.err
        assert "element 0 (text): bbox" in captured.err, captured.err
        assert "missing or unreadable" in captured.err, captured.err
        assert "skipped" not in captured.out, "stdout stays clean"
        assert "1 element (shape: 1)" in captured.out, captured.out
        assert len(_assert_real_deck(out).slides[0].shapes) == 1

    @pytest.mark.parametrize("bad_bbox", [None, "top left"], ids=["null", "string"])
    def test_a_non_container_bbox_behaves_exactly_as_it_always_did(
        self, tmp_path, capsys, bad_bbox
    ):
        """These were already reported; consistency is the point of the fix."""
        code, _out = _from_json(
            tmp_path,
            {"elements": [{"type": "text", "bbox": bad_bbox, "content": "x"}, _GOOD_SHAPE]},
        )
        assert code == 0
        assert "element 0 (text): bbox" in capsys.readouterr().err

    def test_a_slide_of_nothing_but_slivers_is_a_loud_total_failure(self, tmp_path, capsys):
        code, out = _from_json(
            tmp_path, {"elements": [{"type": "text", "bbox": [5, 5], "content": "x"}]}
        )
        assert code == 1, "losing everything may not be reported as success"
        assert "missing or unreadable" in capsys.readouterr().err
        assert not out.exists(), "no deck may be written"

    def test_a_good_spec_is_untouched_by_the_tightening(self, tmp_path, capsys):
        code, out = _from_json(tmp_path, mock_data.SAMPLE_SPEC_JSON)
        captured = capsys.readouterr()
        assert code == 0
        assert "skipped" not in captured.err, captured.err
        assert "below the" not in captured.err, "no bbox in a good spec is a sliver"
        assert len(_assert_real_deck(out).slides[0].shapes) == len(
            mock_data.sample_spec().elements
        )


class TestABareArrayKeepsItsGoodElements:
    """One unknown ``type`` in a bare array destroyed the whole conversion.

    Exit 1, no deck, a raw ``model_type`` ValidationError, no per-element
    reasons, and "That file is not a valid slide spec" - while the identical
    content wrapped in ``{"elements": [...]}`` exited 0 and named the bad one.
    """

    PAYLOAD = [
        {"type": "text", "bbox": [1, 1, 100, 50], "content": "good one"},
        {"type": "hologram", "bbox": [1, 1, 9, 9]},
        {"type": "shape", "bbox": [2, 2, 100, 50], "kind": "circle", "fill_color": "#FF0000"},
    ]

    def test_it_builds_the_two_good_elements_and_names_the_bad_one(self, tmp_path, capsys):
        code, out = _from_json(tmp_path, self.PAYLOAD)
        captured = capsys.readouterr()

        assert code == 0
        assert "1 element skipped" in captured.err, captured.err
        assert "hologram" in captured.err, captured.err
        assert "not a valid slide spec" not in captured.err
        assert "2 elements (shape: 1, text: 1)" in captured.out, captured.out
        assert len(_assert_real_deck(out).slides[0].shapes) == 2

    def test_the_bare_and_wrapped_spellings_now_agree(self, tmp_path, capsys):
        bare_code, bare_out = _from_json(tmp_path, self.PAYLOAD, "bare.json")
        bare = capsys.readouterr()
        wrapped_code, wrapped_out = _from_json(
            tmp_path, {"elements": self.PAYLOAD}, "wrapped.json"
        )
        wrapped = capsys.readouterr()

        def _summary(text: str) -> list[str]:
            # everything but the "Wrote <path>" line, whose path differs by design
            return [ln for ln in text.splitlines() if not ln.startswith("Wrote ")]

        assert bare_code == wrapped_code == 0
        assert _summary(bare.out) == _summary(wrapped.out)
        assert len(_assert_real_deck(bare_out).slides[0].shapes) == len(
            _assert_real_deck(wrapped_out).slides[0].shapes
        )

    @pytest.mark.parametrize("payload", [[1, 2, 3], ["a", "b"], [], [{"note": "hi"}]])
    def test_an_array_that_is_not_elements_at_all_is_still_a_bad_spec(
        self, tmp_path, capsys, payload
    ):
        """Recovery must not start swallowing lists that are something else."""
        code, out = _from_json(tmp_path, payload)
        assert code == 1
        assert "not a valid slide spec" in capsys.readouterr().err
        assert not out.exists()


class TestDroppedCannotBeForgedByTheInput:
    """A ``dropped`` key in the payload was printed as our own diagnostic.

    Exit 0 plus "1 element skipped ... element 0 (text): TOTALLY MADE UP" -
    verbatim from the input, with nothing actually dropped.
    """

    FORGED = ["element 0 (text): TOTALLY MADE UP"]

    def test_the_reproduction_no_longer_invents_a_drop(self, tmp_path, capsys):
        code, out = _from_json(
            tmp_path, {"background_color": "#FFFFFF", "dropped": list(self.FORGED)}
        )
        captured = capsys.readouterr()

        assert code == 0 and out.exists()
        assert "TOTALLY MADE UP" not in captured.err, captured.err
        assert "skipped" not in captured.err, captured.err
        # the spec really is empty, and THAT is what must be reported
        assert "blank" in captured.err.lower()

    def test_a_forged_drop_cannot_ride_along_with_a_real_one(self, tmp_path, capsys):
        code, _out = _from_json(
            tmp_path,
            {"elements": [_GOOD_TEXT, _TABLE_NO_DATA], "dropped": list(self.FORGED)},
        )
        err = capsys.readouterr().err
        assert code == 0
        assert "1 element skipped" in err, "the real drop is still reported"
        assert "element 1 (table): data - Field required" in err
        assert "TOTALLY MADE UP" not in err, err

    def test_a_forged_drop_in_the_models_own_reply_is_ignored_too(
        self, tmp_path, monkeypatch, capsys
    ):
        """The same injection is reachable through the vision path."""
        reply = json.dumps(
            {"elements": [_GOOD_TEXT], "dropped": ["element 9 (text): FROM THE MODEL"]}
        )
        monkeypatch.setattr(config, "get_client", lambda: _ReplyClient(reply), raising=False)
        png = _sample_png(tmp_path)
        out = tmp_path / "injected.pptx"

        assert _run(["-i", str(png), "-o", str(out)]) == 0
        captured = capsys.readouterr()
        assert "FROM THE MODEL" not in captured.err, captured.err
        assert "skipped" not in captured.err, captured.err
        assert len(_assert_real_deck(out).slides[0].shapes) == 1

    def test_save_json_still_never_writes_the_key(self, tmp_path, capsys):
        spec_path = tmp_path / "spec.json"
        spec_path.write_text(
            json.dumps({"elements": [_GOOD_TEXT], "dropped": list(self.FORGED)}),
            encoding="utf-8",
        )
        dump = tmp_path / "dump.json"
        assert (
            _run(
                [
                    "--from-json",
                    str(spec_path),
                    "-o",
                    str(tmp_path / "o.pptx"),
                    "--save-json",
                    str(dump),
                ]
            )
            == 0
        )
        capsys.readouterr()
        assert "dropped" not in dump.read_text(encoding="utf-8")


class TestTheDeckDeclaresItsRealSlideSize:
    def test_a_built_package_says_screen16x9_not_screen4x3(self, tmp_path, capsys):
        out = tmp_path / "sized.pptx"
        assert _run(["--from-json", str(_spec_file(tmp_path)), "-o", str(out)]) == 0
        capsys.readouterr()

        with zipfile.ZipFile(out) as zf:
            xml = zf.read("ppt/presentation.xml").decode("utf-8")
        declared = re.search(r"<p:sldSz[^>]*>", xml)
        assert declared is not None, xml[:400]
        assert 'type="screen16x9"' in declared.group(0), declared.group(0)
        assert "screen4x3" not in xml


# ---------------------------------------------------------------------------
# usage errors
# ---------------------------------------------------------------------------
class TestUsageErrors:
    def test_unsupported_extension_returns_two(self, tmp_path, capsys):
        bad = tmp_path / "screenshot.bmp"
        bad.write_bytes(b"\x00\x01\x02")
        assert _run(["-i", str(bad), "-o", str(tmp_path / "x.pptx")]) == 2
        assert capsys.readouterr().err.strip()

    def test_missing_input_file_returns_two(self, tmp_path):
        assert _run(["-i", str(tmp_path / "ghost.png"), "-o", str(tmp_path / "x.pptx")]) == 2

    def test_no_arguments_at_all_returns_two(self, capsys):
        assert _run([]) == 2
        capsys.readouterr()

    def test_help_exits_zero(self, capsys):
        assert _run(["--help"]) == 0
        assert "usage" in capsys.readouterr().out.lower()


# ---------------------------------------------------------------------------
# make_sample_image
# ---------------------------------------------------------------------------
class TestMakeSampleImage:
    def test_default_size_is_1600x900_16_by_9(self, tmp_path):
        path = make_sample_image(tmp_path / "s.png")
        assert path.exists() and path.is_absolute()
        with Image.open(path) as img:
            assert img.size == (1600, 900)
            assert img.mode == "RGB"
            assert img.format == "PNG"

    def test_custom_size_is_respected(self, tmp_path):
        path = make_sample_image(tmp_path / "small.png", size=(640, 360))
        with Image.open(path) as img:
            assert img.size == (640, 360)

    def test_accepts_a_string_path_and_creates_parents(self, tmp_path):
        path = make_sample_image(str(tmp_path / "nested" / "s.png"), size=(320, 180))
        assert Path(path).exists()

    def test_the_image_is_not_blank(self, tmp_path):
        path = make_sample_image(tmp_path / "s.png", size=(800, 450))
        with Image.open(path) as img:
            colors = img.convert("RGB").getcolors(maxcolors=1 << 20)
        assert colors is not None
        assert len(colors) > 8, "the sample slide looks blank"

    def test_the_header_band_and_body_are_distinguishable(self, tmp_path):
        """A crude layout check: dark band on top, light background below."""
        path = make_sample_image(tmp_path / "s.png")
        with Image.open(path) as img:
            rgb = img.convert("RGB")
            band = rgb.getpixel((20, 20))
            body = rgb.getpixel((20, 860))
        assert sum(band) < sum(body), "the header band should be darker than the body"
        assert band != (255, 255, 255)

    def test_default_filename_lands_in_the_cwd(self, tmp_path, monkeypatch):
        monkeypatch.chdir(tmp_path)
        path = make_sample_image()
        assert path == (tmp_path / "sample.png").resolve()
        assert path.exists()

    def test_output_is_deterministic(self, tmp_path):
        a = make_sample_image(tmp_path / "a.png", size=(400, 225))
        b = make_sample_image(tmp_path / "b.png", size=(400, 225))
        assert Path(a).read_bytes() == Path(b).read_bytes()

    def test_rejects_a_degenerate_size(self, tmp_path):
        with pytest.raises(ValueError):
            make_sample_image(tmp_path / "bad.png", size=(0, 100))


# ---------------------------------------------------------------------------
# cross-module: generated image -> stubbed LLM -> native deck
# ---------------------------------------------------------------------------
def test_full_pipeline_on_a_generated_image_yields_no_pictures(tmp_path, monkeypatch):
    """The generated slide goes in; an editable deck with no images comes out."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    png = make_sample_image(tmp_path / "sample.png", size=(800, 450))

    def fake_analyze(
        image_path, model=None, client=None, extra_instructions=None, verbose=False
    ):
        assert Path(str(image_path)).exists()
        return mock_data.sample_spec()

    _patch_analyzer(monkeypatch, fake_analyze)
    out = tmp_path / "pipeline.pptx"
    assert _run(["-i", str(png), "-o", str(out)]) == 0

    slide = Presentation(str(out)).slides[0]
    assert all(s.shape_type != MSO_SHAPE_TYPE.PICTURE for s in slide.shapes)
    assert "<p:pic" not in slide.shapes._spTree.xml

    kinds = {
        "table": sum(1 for s in slide.shapes if s.has_table),
        "text": sum(1 for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX),
        "shape": sum(1 for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE),
    }
    spec = mock_data.sample_spec()
    assert kinds["table"] == sum(1 for el in spec.elements if isinstance(el, TableElement))
    assert kinds["text"] == sum(1 for el in spec.elements if isinstance(el, TextElement))
    # ">=": the slide background may legitimately cost one extra autoshape.
    assert kinds["shape"] >= sum(1 for el in spec.elements if isinstance(el, ShapeElement))


# =========================================================================== #
# the README's own keyless recipes
# =========================================================================== #
README = Path(__file__).resolve().parents[2] / "README.md"

#: A promise that a command needs no proxy / key / network.
KEYLESS_CLAIM = re.compile(r"no\s+(?:proxy|key|network)", re.I)

#: An invocation that calls the vision model, i.e. the ONE CLI path that needs
#: ``LITELLM_API_KEY`` and a reachable proxy.
VISION_CALL = re.compile(r"cli\.py\b[^\n]*(?:\s-i\s|\s--input\s)")

readme_exists = pytest.mark.skipif(
    not README.is_file(), reason="README.md is not present in this checkout"
)


def _bash_blocks(text: str) -> list[tuple[list[str], list[str]]]:
    """``(the paragraph immediately above, the block's lines)`` per bash fence.

    The "paragraph immediately above" is the run of non-empty lines touching
    the fence - that is the sentence a reader takes as describing the command,
    and it is where "no proxy required" lived. Prose further up is not
    considered, so an unrelated mention of "no key" three paragraphs earlier
    cannot govern a command.
    """
    lines = text.splitlines()
    blocks: list[tuple[list[str], list[str]]] = []
    index = 0
    while index < len(lines):
        if lines[index].strip().startswith("```bash"):
            start = index + 1
            end = start
            while end < len(lines) and not lines[end].strip().startswith("```"):
                end += 1
            cursor = start - 2
            while cursor >= 0 and not lines[cursor].strip():  # the blank line(s)
                cursor -= 1
            lead: list[str] = []
            while cursor >= 0 and lines[cursor].strip():
                lead.insert(0, lines[cursor])
                cursor -= 1
            blocks.append((lead, lines[start:end]))
            index = end
        index += 1
    return blocks


def _keyless_commands(text: str | None = None) -> list[tuple[str, str]]:
    """``(the claim, the command)`` for every command a keyless promise covers.

    Inside a block, a RUN of ``#`` comment lines governs the commands beneath
    it (a run, not the last line: example 5's promise is on the first of six
    comment lines, and taking only the nearest one read the promise off the
    wrong sentence). Before the first comment, the lead paragraph governs.
    """
    if text is None:
        text = README.read_text(encoding="utf-8")
    governed: list[tuple[str, str]] = []
    for lead, body in _bash_blocks(text):
        lead_claim = "\n".join(lead)
        comments: list[str] = []
        start_new = True
        for line in body:
            stripped = line.strip()
            if not stripped:
                start_new = True
                continue
            if stripped.startswith("#"):
                if start_new:
                    comments = [stripped]
                    start_new = False
                else:
                    comments.append(stripped)
                continue
            claim = "\n".join(comments) if comments else lead_claim
            if KEYLESS_CLAIM.search(claim):
                governed.append((claim, stripped))
            start_new = True
    return governed


@readme_exists
class TestTheReadmesKeylessRecipesReallyAreKeyless:
    """Two recipes promised keylessness and then ran a vision conversion.

    When ``--mock`` was removed from the docs the flag was swapped for
    ``-i sample.png`` but the promise above it was left alone, so "Quick manual
    smoke test, no proxy required" and "# 5. no proxy, no key, no network" both
    became false - and because a failed analysis writes no file, the ``&& open``
    half never ran and example 5's second line then failed too ("JSON spec not
    found"). This is the first thing a new contributor runs.
    """

    def test_the_scan_finds_the_recipes(self):
        governed = _keyless_commands()
        assert governed, "no keyless promise found in the README - re-point this"
        assert any("--from-json" in command for _, command in governed)

    def test_no_keyless_promise_covers_a_vision_command(self):
        offenders = [
            f"{claim.strip()!r} -> {command!r}"
            for claim, command in _keyless_commands()
            if VISION_CALL.search(command)
        ]
        assert offenders == [], (
            "a command that needs LITELLM_API_KEY and a reachable proxy sits "
            f"under a promise that it needs neither: {offenders}"
        )

    def test_the_guard_can_actually_fail(self):
        assert VISION_CALL.search("$PY cli.py -i sample.png -o /tmp/first.pptx")
        assert VISION_CALL.search("../.venv/bin/python cli.py -i sample.png -o x.pptx")
        assert not VISION_CALL.search("$PY cli.py --from-json /tmp/spec.json -o x.pptx")
        assert KEYLESS_CLAIM.search("Quick manual smoke test, no proxy required:")
        assert KEYLESS_CLAIM.search("# 5. no proxy, no key, no network: prove it")
        assert not KEYLESS_CLAIM.search("# 1. the happy path")

    def test_every_keyless_command_is_one_the_cli_can_really_run_offline(self):
        """Each keyless line must either build from a spec or produce one - the
        recipe has to be self-contained, or step 2 fails on the file step 1
        never wrote."""
        for claim, command in _keyless_commands():
            if "cli.py" not in command:
                continue
            assert "--from-json" in command, (
                f"{command!r} is offered as keyless but is not the --from-json "
                f"path (claim: {claim.strip()!r})"
            )

    def test_the_smoke_test_recipe_runs_end_to_end_with_no_key(self, tmp_path, monkeypatch):
        """Not a source assertion: this actually performs what the README's
        smoke test does - seed a spec from the fixture, then build from it with
        no key in the environment at all."""
        monkeypatch.delenv("LITELLM_API_KEY", raising=False)
        monkeypatch.setattr(config, "LITELLM_API_KEY", "", raising=False)

        spec_path = tmp_path / "spec.json"
        spec_path.write_text(mock_data.SAMPLE_SPEC_JSON, encoding="utf-8")
        out = tmp_path / "demo.pptx"

        assert _run(["--from-json", str(spec_path), "-o", str(out)]) == 0
        assert out.is_file(), "the README's smoke test wrote no deck"
        assert len(Presentation(str(out)).slides) == 1

    def test_the_guard_catches_the_text_it_was_written_for(self):
        """The exact two recipes that shipped broken, fed through the scanner
        without touching the real README - a guard nobody has seen fail is a
        guard nobody knows works."""
        broken = "\n".join(
            [
                "Quick manual smoke test, no proxy required:",
                "",
                "```bash",
                "cd backend && ../.venv/bin/python cli.py -i sample.png "
                "-o /tmp/demo.pptx && open /tmp/demo.pptx",
                "```",
                "",
                "Examples:",
                "",
                "```bash",
                "# 5. no proxy, no key, no network: prove the builder works "
                "from a saved spec",
                "$PY cli.py -i sample.png --save-json /tmp/spec.json "
                "-o /tmp/first.pptx",
                "$PY cli.py --from-json /tmp/spec.json -o /tmp/demo.pptx",
                "```",
            ]
        )
        flagged = [
            command
            for _, command in _keyless_commands(broken)
            if VISION_CALL.search(command)
        ]
        assert len(flagged) == 2, (
            "the scanner no longer reports the two commands this guard exists "
            f"for: {flagged}"
        )
        # ...and the honest half of the old example is NOT flagged.
        assert all("--save-json" in c or "-o /tmp/demo.pptx &&" in c for c in flagged)
