"""Tests for the HTTP layer: ``app.main`` / ``app.api.routes`` / the converter.

Everything here is **fully offline**. Two rules keep it that way and they are
asserted rather than assumed:

* the conversion path is exercised by monkeypatching
  ``vision_analyzer.analyze_image`` (the module attribute, which the converter
  looks up at call time), never by letting a request reach the proxy - see the
  ``offline_convert`` fixture, which is how every test here gets a 200 out of
  ``/api/convert``;
* nothing may build a real client, so ``config.get_client`` is replaced with a
  bomb in the fixtures that cover it.

There is no mock mode to lean on: it was removed from the product, and
``TestMockModeIsGone`` asserts that it stays removed. ``app.mock_data`` is still
imported below, but only as this suite's shared spec FIXTURE.

The API key is likewise never real: ``config.LITELLM_API_KEY`` is a module
global read at call time, so a fake value is patched onto the module instead of
touching the environment - and the fake is an obvious placeholder, because
``tests/test_config.py`` scans this file for the real key's literal.
"""

from __future__ import annotations

import base64
import copy
import io
import json
import pathlib
import re
import time
import urllib.parse
import uuid
import zipfile

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app import config, mock_data, pptx_builder, vision_analyzer
from app.main import ALLOWED_ORIGINS, app, error_body
from app.schemas import SlideSpec
from app.services import converter
from app.services.converter import ConversionFailure, DownloadStore

from make_sample_image import make_sample_image

#: An obviously fake key. Never the real one (see the module docstring).
FAKE_KEY = "test-key-not-a-real-credential"

PPTX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)


# --------------------------------------------------------------------------- #
# fixtures
# --------------------------------------------------------------------------- #
@pytest.fixture
def client() -> TestClient:
    """A client that surfaces a crash as a 500 response instead of re-raising,
    so the app's unexpected-exception handler is actually observable."""
    with TestClient(app, raise_server_exceptions=False) as test_client:
        yield test_client


@pytest.fixture(autouse=True)
def _isolated_download_store():
    """Each test gets an empty, process-local download store."""
    converter.downloads.clear()
    yield
    converter.downloads.clear()


@pytest.fixture(autouse=True)
def _no_real_key(monkeypatch):
    """Default to "no key configured" so nothing can accidentally use one."""
    monkeypatch.setattr(config, "LITELLM_API_KEY", "", raising=False)


@pytest.fixture
def with_key(monkeypatch):
    monkeypatch.setattr(config, "LITELLM_API_KEY", FAKE_KEY, raising=False)
    return FAKE_KEY


@pytest.fixture
def no_client_allowed(monkeypatch):
    """Make building an LLM client an error: proves a path is truly offline."""

    def _bomb(*args, **kwargs):  # pragma: no cover - a call is the failure
        raise AssertionError("config.get_client() must not be called on this path")

    monkeypatch.setattr(config, "get_client", _bomb)


@pytest.fixture(scope="module")
def png_bytes(tmp_path_factory) -> bytes:
    """A real, decodable PNG (the project's own sample generator)."""
    path = tmp_path_factory.mktemp("images") / "sample.png"
    make_sample_image(str(path))
    return path.read_bytes()


def upload(data: bytes, name: str = "sample.png", content_type: str = "image/png"):
    return {"file": (name, data, content_type)}


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
def envelope(response) -> dict:
    """The error envelope, asserted to have exactly the contract's shape."""
    body = response.json()
    assert set(body) == {"detail"}, body
    detail = body["detail"]
    assert isinstance(detail, dict), detail
    assert set(detail) == {"code", "message", "hint"}, detail
    for key, value in detail.items():
        assert isinstance(value, str), (key, value)
    assert detail["code"] and detail["message"]
    return detail


def open_deck(data: bytes) -> Presentation:
    assert data[:2] == b"PK", "not a zip container, so not a .pptx"
    return Presentation(io.BytesIO(data))


def assert_native_deck(data: bytes) -> Presentation:
    """python-pptx can open it, it is 16:9, and it is NOT a flattened slide.

    This used to assert "zero picture shapes", which stopped being the right
    rule the moment a genuine photo / chart / logo region became a legitimate
    ``image`` element. The replacement is ``pptx_builder.flattening_report`` -
    THE definition of "flattened" for this project, so the tests and the build
    itself read the same one. A second copy of the rule living here is exactly
    how the two would drift and how a flattened deck would end up shipped past
    whichever of them happened to be laxer.
    """
    presentation = open_deck(data)
    assert len(presentation.slides) == 1
    assert round(presentation.slide_width.inches, 3) == 13.333
    assert round(presentation.slide_height.inches, 3) == 7.5

    report = pptx_builder.flattening_report(presentation)
    assert report["flattened"] is False, report
    assert report["violations"] == [], report
    # A picture whose bytes cannot be resolved is a deck PowerPoint offers to
    # repair, and every OTHER number here reads normally on one - shape counts,
    # geometry, even ``len(media) <= picture_count`` (media is 0). Without this
    # line the helper blessed a corrupt file.
    assert report["unreadable_picture_count"] == 0, report

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = archive.namelist()
    media = [name for name in names if name.startswith("ppt/media/")]
    # Media is allowed now, but ONLY as much of it as there are pictures: a
    # deck carrying image parts that no shape references would mean the crop
    # went into the package without reaching the slide. Identical crops share
    # ONE part, hence "<=" and not "==".
    assert len(media) <= report["picture_count"]
    if report["picture_count"]:
        assert media, "a picture is on the slide but its bytes are not in the package"
    assert _every_embed_resolves(data)
    assert any(name.startswith("ppt/slides/slide") for name in names)
    return presentation


def _every_embed_resolves(data: bytes) -> bool:
    """Does every ``r:embed`` in every slide have a matching relationship?

    A dangling one is an OPC violation - the "found a problem with content"
    repair prompt - and it is invisible to python-pptx's own reader, so it has
    to be checked against the raw package.
    """
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        slides = [
            name
            for name in archive.namelist()
            if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
        ]
        for name in slides:
            xml = archive.read(name).decode("utf-8")
            rels_name = name.replace("slides/", "slides/_rels/") + ".rels"
            rels = (
                archive.read(rels_name).decode("utf-8")
                if rels_name in archive.namelist()
                else ""
            )
            declared = set(re.findall(r'Id="([^"]+)"', rels))
            dangling = [
                rid
                for rid in re.findall(r'r:embed="([^"]+)"', xml)
                if rid not in declared
            ]
            assert dangling == [], f"{name} references {dangling} with no relationship"
    return True


def spec_dict() -> dict:
    return copy.deepcopy(mock_data.sample_spec_dict())


def fake_analyzer(spec: SlideSpec | None = None, *, stderr: str = ""):
    """A stand-in for ``analyze_image`` that records how it was called."""
    calls: list[dict] = []

    def _analyze(image_path, model=None, client=None, extra_instructions=None, verbose=False):
        calls.append(
            {
                "image_path": str(image_path),
                "model": model,
                "client": client,
                "extra_instructions": extra_instructions,
                "bytes": len(open(image_path, "rb").read()),
            }
        )
        if stderr:
            import sys

            print(stderr, file=sys.stderr)
        return spec if spec is not None else mock_data.sample_spec()

    _analyze.calls = calls  # type: ignore[attr-defined]
    return _analyze


@pytest.fixture
def offline_convert(monkeypatch, with_key, no_client_allowed):
    """The real ``/api/convert`` path, offline.

    A (fake) key is configured so ``require_api_key`` passes, ``analyze_image``
    is a stand-in that returns the sample spec, and building a real client is
    still an outright failure - so a 200 here proves the whole pipeline ran
    without the network, which is what mock mode used to be used for.

    Returns the stand-in, so a test can also assert how it was called.
    """
    analyze = fake_analyzer()
    monkeypatch.setattr(vision_analyzer, "analyze_image", analyze)
    return analyze


# =========================================================================== #
# GET /api/health
# =========================================================================== #
class TestHealth:
    def test_shape_is_exactly_the_contract(self, client):
        response = client.get("/api/health")
        assert response.status_code == 200
        body = response.json()

        assert set(body) == {"status", "proxy", "supported_extensions", "slide"}
        assert body["status"] == "ok"
        assert set(body["proxy"]) == {"base_url", "model", "api_key_configured"}
        assert set(body["slide"]) == {"width_in", "height_in", "coord_scale"}

        assert body["proxy"]["base_url"] == config.LITELLM_BASE_URL
        assert body["proxy"]["model"] == config.DEFAULT_MODEL
        assert body["supported_extensions"] == [".jpeg", ".jpg", ".png", ".webp"]
        assert body["slide"] == {
            "width_in": 13.333,
            "height_in": 7.5,
            "coord_scale": 1000,
        }

    def test_api_key_configured_is_a_bool_not_the_key(self, client, with_key):
        body = client.get("/api/health").json()
        flag = body["proxy"]["api_key_configured"]
        assert flag is True
        assert not isinstance(flag, str)

    def test_it_reports_false_when_no_key_is_set(self, client):
        assert client.get("/api/health").json()["proxy"]["api_key_configured"] is False

    @pytest.mark.parametrize("key", [FAKE_KEY, "sk-abcdef0123456789", "x" * 64])
    def test_the_key_never_appears_anywhere_in_the_response(
        self, client, monkeypatch, key
    ):
        """Not the key, not a prefix of it, not its length."""
        monkeypatch.setattr(config, "LITELLM_API_KEY", key, raising=False)
        response = client.get("/api/health")
        text = response.text

        assert key not in text
        for length in range(4, min(len(key), 12) + 1):
            assert key[:length] not in text, f"leaked a {length}-char prefix"
        assert str(len(key)) not in json.dumps(response.json()["proxy"])

    def test_it_stays_200_when_the_proxy_is_unreachable(
        self, client, with_key, no_client_allowed
    ):
        """Configuration only: health never makes a network call, so a dead
        proxy cannot take the health check down with it."""
        assert client.get("/api/health").status_code == 200


# =========================================================================== #
# POST /api/convert - the happy path
# =========================================================================== #
class TestConvertHappyPath:
    def test_end_to_end(self, client, png_bytes, offline_convert):
        response = client.post("/api/convert", files=upload(png_bytes))
        assert response.status_code == 200
        body = response.json()

        assert set(body) == {
            "status",
            "counts",
            "total",
            "dropped",
            "warnings",
            "spec",
            "preview",
            "download_id",
            "filename",
        }
        assert body["status"] == "success"
        assert body["counts"] == {"shape": 4, "table": 1, "text": 4}
        assert body["total"] == 9
        assert body["total"] == sum(body["counts"].values())
        assert body["dropped"] == []
        assert body["warnings"] == []
        assert body["filename"] == "sample.pptx"
        assert converter.is_download_id(body["download_id"])
        assert uuid.UUID(hex=body["download_id"]).version == 4

    def test_the_response_carries_a_preview_of_the_deck(
        self, client, png_bytes, offline_convert
    ):
        """The contract's new field. Its contents are pinned in
        ``tests/test_preview.py``; here it only has to arrive, in shape, and to
        describe the same slide the deck does."""
        body = client.post("/api/convert", files=upload(png_bytes)).json()
        preview = body["preview"]

        assert isinstance(preview, dict)
        assert set(preview) == {"slide", "background_color", "shapes"}
        assert preview["slide"]["width_in"] == 13.333
        assert preview["slide"]["height_pt"] == 540.0
        assert preview["background_color"] == "#FFFFFF"
        # One preview entry per shape the deck really holds.
        deck = client.get(f"/api/download/{body['download_id']}")
        presentation = open_deck(deck.content)
        assert len(preview["shapes"]) == len(presentation.slides[0].shapes)

    def test_dropped_and_warnings_are_always_arrays(self, client, png_bytes, offline_convert):
        body = client.post("/api/convert", files=upload(png_bytes)).json()
        assert isinstance(body["dropped"], list)
        assert isinstance(body["warnings"], list)

    def test_the_spec_is_the_validated_slidespec_without_dropped(
        self, client, png_bytes, offline_convert
    ):
        """``SlideSpec.dropped`` is ``exclude=True`` and tool-generated: it is
        reported in its own field and must never appear inside ``spec``."""
        body = client.post("/api/convert", files=upload(png_bytes)).json()
        spec = body["spec"]

        assert "dropped" not in spec
        assert set(spec) == {"background_color", "elements"}
        assert spec["background_color"] == "#FFFFFF"
        assert len(spec["elements"]) == body["total"]
        # It round-trips: the reported spec rebuilds the very same deck.
        assert SlideSpec.model_validate(spec).counts_by_type() == {
            "shape": 4,
            "table": 1,
            "text": 4,
        }

    def test_the_returned_deck_is_native_and_not_flattened(
        self, client, png_bytes, offline_convert
    ):
        """The sample spec has no ``image`` element, so this deck really does
        carry zero pictures - asserted here, on a spec where it is true, rather
        than as a blanket rule (see ``assert_native_deck``)."""
        body = client.post("/api/convert", files=upload(png_bytes)).json()
        deck = client.get(f"/api/download/{body['download_id']}")
        assert deck.status_code == 200

        presentation = assert_native_deck(deck.content)
        kinds = {shape.shape_type for shape in presentation.slides[0].shapes}
        assert kinds <= {
            MSO_SHAPE_TYPE.AUTO_SHAPE,
            MSO_SHAPE_TYPE.TABLE,
            MSO_SHAPE_TYPE.TEXT_BOX,
        }
        assert pptx_builder.flattening_report(presentation)["picture_count"] == 0

    @pytest.mark.parametrize(
        ("uploaded", "expected"),
        [
            ("sample.png", "sample.pptx"),
            ("my slide.jpg", "my slide.pptx"),
            ("../../etc/passwd.png", "passwd.pptx"),
            ("C:\\Users\\bob\\deck.png", "deck.pptx"),
        ],
    )
    def test_the_filename_is_sanitized_from_the_upload_stem(
        self, client, png_bytes, offline_convert, uploaded, expected
    ):
        response = client.post("/api/convert", files=upload(png_bytes, name=uploaded))
        assert response.status_code == 200
        assert response.json()["filename"] == expected

        disposition = client.get(
            f"/api/download/{response.json()['download_id']}"
        ).headers["content-disposition"]
        assert disposition == f'attachment; filename="{expected}"'
        assert "\r" not in disposition and "\n" not in disposition

    def test_a_hostile_upload_name_cannot_inject_a_header(
        self, client, png_bytes, offline_convert
    ):
        """The multipart encoder escapes most of this on the way in; whatever
        survives must still be safe to interpolate into Content-Disposition.

        The JSON ``filename`` is the DISPLAY name (it is UTF-8 JSON and ends up
        in an ``<a download>``, so it keeps the user's own letters), and the
        header's ``filename=`` is the latin-1 form ``pptx_headers`` narrows it
        to. Both have to be free of anything structural, which is what this
        asserts - on the value the server actually produced, since the
        multipart encoder percent-escapes some of this on the way in and the
        server never sees the raw python string."""
        hostile = 'ev"il\r\nX; q=1.png'
        response = client.post("/api/convert", files=upload(png_bytes, name=hostile))
        assert response.status_code == 200
        filename = response.json()["filename"]
        assert filename.endswith(".pptx")
        assert not set(filename) & set('"\r\n;/\\')

        disposition = client.get(
            f"/api/download/{response.json()['download_id']}"
        ).headers["content-disposition"]
        disposition.encode("latin-1")
        assert "\r" not in disposition and "\n" not in disposition
        assert disposition.count('"') == 2
        assert disposition.count(";") == 1  # only the one separating the params
        # This name is pure ASCII, so there is no second parameter to hide in.
        assert "filename*" not in disposition

        header_name = disposition.split('filename="', 1)[1].removesuffix('"')
        assert header_name.endswith(".pptx")
        assert header_name.isascii()
        assert not set(header_name) & set('"\r\n;/\\')
        # ...and it is exactly the documented derivation: the latin-1 narrowing
        # of the display name, not some third value.
        assert header_name == converter.safe_pptx_filename(filename)

        # Independently, the sanitizer handles the raw string too - the
        # encoder's escaping must not be what is keeping us safe.
        for raw in (hostile, 'ev"il\r\nX; q=1'):
            assert not set(converter.safe_pptx_filename(raw)) & set('"\r\n;/\\')
            assert not set(converter.display_pptx_filename(raw)) & set('"\r\n;/\\')


# =========================================================================== #
# Mock mode is gone from the product - and has to stay gone
# =========================================================================== #
class TestMockModeIsGone:
    """``is_mock`` was a real API field; it is not one any more.

    The risk being covered is a silent resurrection: a stale client (or a stale
    README) still sending ``is_mock=true`` must not get a deck built from the
    sample fixture - it must be treated exactly as if it had sent nothing,
    because an unknown multipart field simply is nothing.
    """

    @pytest.mark.parametrize("flag", ["true", "True", "TRUE", "1", "yes", "on"])
    def test_posting_is_mock_enables_nothing_and_still_needs_a_key(
        self, client, png_bytes, monkeypatch, flag
    ):
        """No key + is_mock=true is a 503, the same as no key alone. The
        exemption is what was removed, so this is the load-bearing assertion."""

        def _bomb(*args, **kwargs):  # pragma: no cover - a call is the failure
            raise AssertionError("nothing may analyze or render without a key")

        monkeypatch.setattr(vision_analyzer, "analyze_image", _bomb)
        response = client.post(
            "/api/convert", files=upload(png_bytes), data={"is_mock": flag}
        )
        assert response.status_code == 503
        assert envelope(response)["code"] == "missing_api_key"

    def test_is_mock_is_ignored_rather_than_rejected(
        self, client, png_bytes, offline_convert
    ):
        """An unknown form field is not an error: a stale client gets the real
        conversion it would have got anyway, not a 422."""
        response = client.post(
            "/api/convert", files=upload(png_bytes), data={"is_mock": "true"}
        )
        assert response.status_code == 200
        # It really went through the model path, not a built-in spec.
        assert len(offline_convert.calls) == 1

    def test_the_field_is_not_in_the_published_schema(self, client):
        """The form body is published as a ``$ref``, so the referenced component
        is what has to be read - asserting against the ``$ref`` itself would
        pass no matter which fields the endpoint declares."""
        schema = client.get("/openapi.json").json()
        body = schema["paths"]["/api/convert"]["post"]["requestBody"]
        ref = body["content"]["multipart/form-data"]["schema"]["$ref"]
        assert ref.startswith("#/components/schemas/")
        fields = schema["components"]["schemas"][ref.rsplit("/", 1)[-1]]["properties"]

        assert "is_mock" not in fields
        # The fields that stayed are still declared.
        assert {"file", "model", "extra_instructions"} <= set(fields)

    def test_the_converter_takes_no_mock_argument(self):
        import inspect

        parameters = inspect.signature(converter.convert_upload).parameters
        assert "is_mock" not in parameters
        assert {"model", "extra_instructions", "client"} <= set(parameters)

    @staticmethod
    def _product_modules() -> tuple[pathlib.Path, list[pathlib.Path]]:
        """``backend/`` and every module in it a user can reach.

        Rooted at ``backend/``, NOT at ``backend/app/``: ``--mock`` was a *CLI*
        flag, so ``backend/cli.py`` is the single file this guard most needs to
        look at - and a scan rooted at the package excluded it, along with
        ``make_sample_image.py``. A ``--sample`` flag rebuilt on
        ``mock_data.sample_spec()`` in ``cli.py`` passed the whole suite,
        because nothing spelled "mock" and nothing scanned that file.
        """
        root = pathlib.Path(converter.__file__).resolve().parents[2]
        product = [
            path
            for path in root.rglob("*.py")
            if "mock_data" not in path.name
            and "__pycache__" not in path.parts
            and "tests" not in path.parts
        ]
        return root, product

    def test_the_scan_reaches_the_cli_and_the_package(self):
        """A guard that looks in the wrong place passes forever."""
        root, product = self._product_modules()
        assert root.name == "backend", f"scan root moved to {root}"
        relative = {str(path.relative_to(root)) for path in product}
        for expected in (
            "cli.py",  # where --mock lived
            "make_sample_image.py",
            "app/services/converter.py",
            "app/api/routes.py",
            "app/main.py",
        ):
            assert expected in relative, f"{expected} is no longer scanned"
        assert not any(
            "tests" in path.parts for path in product
        ), "the fixture's own test users must stay out of the product scan"

    def test_no_product_module_imports_the_fixture(self):
        """``app.mock_data`` stays as a TEST fixture. Nothing a user can reach
        may import it - that is the difference between a fixture and a
        feature."""
        root, product = self._product_modules()
        assert product, "found no product modules to scan"
        offenders = [
            str(path.relative_to(root))
            for path in product
            if "mock_data" in path.read_text(encoding="utf-8").replace(
                "mock_data.py", ""
            )
        ]
        assert offenders == [], f"product code still imports the fixture: {offenders}"


# =========================================================================== #
# POST /api/convert - the real (analyzed) path, offline
# =========================================================================== #
class TestConvertAnalyzed:
    def test_it_passes_the_upload_model_and_instructions_through(
        self, client, png_bytes, with_key, monkeypatch
    ):
        analyze = fake_analyzer()
        monkeypatch.setattr(vision_analyzer, "analyze_image", analyze)

        response = client.post(
            "/api/convert",
            files=upload(png_bytes, name="deck.png"),
            data={"model": "some-vision-model", "extra_instructions": "be precise"},
        )
        assert response.status_code == 200
        assert len(analyze.calls) == 1
        call = analyze.calls[0]
        assert call["model"] == "some-vision-model"
        assert call["extra_instructions"] == "be precise"
        # The analyzer takes a path, so the upload is spilled to a temp file
        # with the right suffix - and the bytes it sees are the ones uploaded.
        assert call["image_path"].endswith(".png")
        assert call["bytes"] == len(png_bytes)

    def test_blank_optional_fields_are_treated_as_absent(
        self, client, png_bytes, with_key, monkeypatch
    ):
        analyze = fake_analyzer()
        monkeypatch.setattr(vision_analyzer, "analyze_image", analyze)

        client.post(
            "/api/convert",
            files=upload(png_bytes),
            data={"model": "   ", "extra_instructions": ""},
        )
        assert analyze.calls[0]["model"] is None
        assert analyze.calls[0]["extra_instructions"] is None

    def test_an_omitted_model_field_uses_the_server_default(
        self, client, png_bytes, with_key, monkeypatch
    ):
        """The UI no longer has a model control, so it sends no ``model`` at
        all. That must keep using ``config.DEFAULT_MODEL`` (from
        ``VISION_MODEL``) rather than becoming a blank model name: the route
        passes ``None`` and the analyzer resolves it - see
        ``test_vision_analyzer.py::test_model_defaults_to_config``."""
        analyze = fake_analyzer()
        monkeypatch.setattr(vision_analyzer, "analyze_image", analyze)

        response = client.post("/api/convert", files=upload(png_bytes))

        assert response.status_code == 200
        assert analyze.calls[0]["model"] is None
        # ...and the configured model is still what the UI can SHOW, so hiding
        # the control does not hide which model ran.
        assert client.get("/api/health").json()["proxy"]["model"] == config.DEFAULT_MODEL

    def test_the_model_form_field_stays_in_the_api(self, client):
        """Hidden on the SCREEN, kept in the contract: API clients and the
        CLI's ``-m`` still need it, so a UI cleanup must not delete it."""
        schema = client.get("/openapi.json").json()
        body = schema["paths"]["/api/convert"]["post"]["requestBody"]
        ref = body["content"]["multipart/form-data"]["schema"]["$ref"]
        fields = schema["components"]["schemas"][ref.rsplit("/", 1)[-1]]["properties"]

        assert "model" in fields
        assert "extra_instructions" in fields
        assert schema["paths"]["/api/convert"]["post"]["requestBody"]["required"]

    def test_the_temp_upload_file_is_removed_afterwards(
        self, client, png_bytes, with_key, monkeypatch
    ):
        analyze = fake_analyzer()
        monkeypatch.setattr(vision_analyzer, "analyze_image", analyze)

        client.post("/api/convert", files=upload(png_bytes))
        import os

        assert not os.path.exists(analyze.calls[0]["image_path"])

    def test_partial_drops_are_reported_and_the_rest_still_builds(
        self, client, png_bytes, with_key, monkeypatch
    ):
        payload = spec_dict()
        payload["elements"].append({"type": "table", "bbox": [10, 10, 100, 100]})
        spec = SlideSpec.model_validate(payload)
        assert spec.dropped, "fixture must actually drop one element"
        monkeypatch.setattr(vision_analyzer, "analyze_image", fake_analyzer(spec))

        response = client.post("/api/convert", files=upload(png_bytes))
        assert response.status_code == 200
        body = response.json()

        assert body["dropped"] == spec.dropped
        assert len(body["dropped"]) == 1
        assert "data - Field required" in body["dropped"][0]
        assert body["total"] == 9
        assert "dropped" not in body["spec"]

    def test_stderr_warnings_are_captured_into_the_warnings_array(
        self, client, png_bytes, with_key, monkeypatch
    ):
        """``pptx_builder`` reports non-fatal problems on stderr; the contract
        requires them as a JSON array, with the tool prefix stripped."""
        thin = {
            "background_color": "#FFFFFF",
            "elements": [
                {
                    "type": "shape",
                    "kind": "rectangle",
                    "bbox": [505, 601, 461, 5],
                    "fill_color": "#123456",
                }
            ],
        }
        monkeypatch.setattr(
            vision_analyzer,
            "analyze_image",
            fake_analyzer(SlideSpec.model_validate(thin)),
        )

        body = client.post("/api/convert", files=upload(png_bytes)).json()
        assert len(body["warnings"]) == 1
        warning = body["warnings"][0]
        assert warning.startswith("bbox (505, 601, 461, 5):")
        assert "below the 0.05in minimum" in warning
        assert "[pptx_builder]" not in warning

    def test_analyzer_warnings_are_captured_with_their_prefix_stripped(
        self, client, png_bytes, with_key, monkeypatch
    ):
        monkeypatch.setattr(
            vision_analyzer,
            "analyze_image",
            fake_analyzer(stderr="[vision_analyzer] model retried without response_format"),
        )
        body = client.post("/api/convert", files=upload(png_bytes)).json()
        assert "model retried without response_format" in body["warnings"]
        assert all("[vision_analyzer]" not in line for line in body["warnings"])

    def test_the_drop_echo_is_not_duplicated_into_warnings(
        self, client, png_bytes, with_key, monkeypatch
    ):
        """``analyze_image`` echoes every drop reason to stderr as well as into
        ``spec.dropped``; the same list must not fill both response fields."""
        payload = spec_dict()
        payload["elements"].append({"type": "table", "bbox": [10, 10, 100, 100]})
        spec = SlideSpec.model_validate(payload)
        reason = spec.dropped[0]
        monkeypatch.setattr(
            vision_analyzer,
            "analyze_image",
            fake_analyzer(spec, stderr=f"[vision_analyzer] dropped {reason}"),
        )

        body = client.post("/api/convert", files=upload(png_bytes)).json()
        assert body["dropped"] == [reason]
        assert body["warnings"] == []

    def test_stderr_is_restored_after_a_request(self, client, png_bytes, offline_convert):
        import sys

        before = sys.stderr
        client.post("/api/convert", files=upload(png_bytes))
        assert sys.stderr is before


# =========================================================================== #
# POST /api/convert - connectors and cropped sub-images, end to end
# =========================================================================== #
#: The region ``DIAGRAM_SPEC``'s image element claims, in normalized units, and
#: the colour painted into exactly that region of the uploaded image. A crop
#: taken from anywhere else is not this colour, so the assertion does not have
#: to trust the builder's arithmetic to check the pipeline's plumbing.
IMAGE_REGION = (600, 200, 300, 400)
IMAGE_RGB = (220, 40, 40)

#: The uploaded image's own pixel size. Deliberately neither square nor the
#: slide's 16:9, so mapping a bbox onto the wrong axis lands on white.
UPLOAD_PX = (1000, 700)


def flat_png(size: tuple[int, int] = (1200, 800)) -> bytes:
    """A uniform upload - a white document scan, a solid-colour deck.

    Two equal-sized crops of it are BYTE-IDENTICAL, which is what makes
    python-pptx de-duplicate them onto one image part with one shared
    relationship id. No duplicate elements in the spec are needed.
    """
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", size, (252, 252, 250)).save(buffer, "PNG")
    return buffer.getvalue()


def diagram_png() -> bytes:
    """A PNG whose :data:`IMAGE_REGION` is solid :data:`IMAGE_RGB`, else white."""
    from PIL import Image

    width, height = UPLOAD_PX
    image = Image.new("RGB", UPLOAD_PX, (255, 255, 255))
    x, y, w, h = IMAGE_REGION
    box = (
        round(x / config.COORD_SCALE * width),
        round(y / config.COORD_SCALE * height),
        round((x + w) / config.COORD_SCALE * width),
        round((y + h) / config.COORD_SCALE * height),
    )
    image.paste(Image.new("RGB", (box[2] - box[0], box[3] - box[1]), IMAGE_RGB), box[:2])
    buffer = io.BytesIO()
    image.save(buffer, "PNG")
    return buffer.getvalue()


#: An architecture-diagram spec: two boxes, a labelled arrow between them, a
#: dashed elbow, a block arrow, and one photo region the analyzer could not
#: express natively. This is the shape of payload the whole feature exists for.
DIAGRAM_SPEC: dict = {
    "background_color": "#FFFFFF",
    "elements": [
        {"type": "shape", "kind": "rectangle", "bbox": [50, 250, 200, 100],
         "fill_color": "#E8EEF7"},
        {"type": "shape", "kind": "rectangle", "bbox": [480, 250, 200, 100],
         "fill_color": "#E8EEF7"},
        {"type": "shape", "kind": "arrow_right", "bbox": [300, 270, 120, 60],
         "fill_color": "#1E3A6E"},
        {"type": "connector", "kind": "straight",
         "start_point": [125, 300], "end_point": [480, 300],
         "arrow_start": False, "arrow_end": True,
         "dash_style": "solid", "color": "#1E3A6E", "width_pt": 1.5},
        {"type": "connector", "kind": "elbow",
         "start_point": [600, 100], "end_point": [900, 600],
         "dash_style": "dashed"},
        {"type": "image", "bbox": list(IMAGE_REGION), "description": "팀 사진"},
        {"type": "text", "bbox": [50, 50, 400, 80], "content": "구성도"},
    ],
}


def diagram_preview(client, image: bytes, monkeypatch) -> dict:
    """Run ``/api/convert`` on ``image`` with a spec containing both new kinds."""
    monkeypatch.setattr(
        vision_analyzer,
        "analyze_image",
        fake_analyzer(SlideSpec.model_validate(copy.deepcopy(DIAGRAM_SPEC))),
    )
    response = client.post("/api/convert", files=upload(image, "diagram.png"))
    assert response.status_code == 200, response.text
    return response.json()


class TestConvertDrawsConnectorsAndCropsSubImages:
    """The end-to-end proof that the uploaded BYTES reach the builder.

    ``/api/convert`` is the only entry point that holds the source image, so it
    is the only one that can crop an ``image`` element. If the bytes were not
    threaded through, this whole class fails with the element skipped and a
    "no source image" warning - which is exactly the failure mode the
    ``/api/preview`` tests assert on purpose.
    """

    @pytest.fixture
    def body(self, client, monkeypatch, with_key, no_client_allowed) -> dict:
        return diagram_preview(client, diagram_png(), monkeypatch)

    def test_the_deck_carries_a_native_connector_and_a_real_picture(
        self, client, body
    ):
        deck = client.get(f"/api/download/{body['download_id']}")
        assert deck.status_code == 200
        presentation = open_deck(deck.content)
        types_seen = [shape.shape_type for shape in presentation.slides[0].shapes]

        # LINE, not CONNECTOR: python-pptx 1.0.2 has no CONNECTOR member, so a
        # test looking for one would pass by never matching anything.
        assert types_seen.count(MSO_SHAPE_TYPE.LINE) == 2
        assert types_seen.count(MSO_SHAPE_TYPE.PICTURE) == 1
        assert MSO_SHAPE_TYPE.AUTO_SHAPE in types_seen
        assert MSO_SHAPE_TYPE.TEXT_BOX in types_seen

    def test_the_deck_is_not_flattened_even_though_it_holds_a_picture(
        self, client, body
    ):
        deck = client.get(f"/api/download/{body['download_id']}")
        presentation = assert_native_deck(deck.content)
        report = pptx_builder.flattening_report(presentation)

        assert report["picture_count"] == 1
        assert report["flattened"] is False
        # 30% x 40% of the slide, well inside the area limit.
        assert report["slides"][0]["picture_area_pct"] == pytest.approx(12.0, abs=0.5)

    def test_a_rejected_duplicate_crop_does_not_corrupt_the_served_deck(
        self, client, monkeypatch, with_key, no_client_allowed
    ):
        """The file the user downloads has to open, not offer to repair.

        Two equal-sized regions of a flat area crop to byte-identical PNGs, and
        python-pptx gives both pictures the SAME image relationship. Their union
        covers 90% of the slide, so the anti-flattening pass drops one - and
        that drop used to delete the relationship the SURVIVOR still pointed at,
        leaving a dangling ``r:embed`` and no ``ppt/media`` part at all. The
        response was still a 200 and the warnings still said the .pptx was fine.
        """
        spec = {
            "elements": [
                {"type": "text", "bbox": [0, 0, 400, 60], "content": "구성도"},
                {"type": "image", "bbox": [0, 50, 500, 900], "description": "왼쪽"},
                {"type": "image", "bbox": [500, 50, 500, 900], "description": "오른쪽"},
            ]
        }
        monkeypatch.setattr(
            vision_analyzer,
            "analyze_image",
            fake_analyzer(SlideSpec.model_validate(spec)),
        )
        response = client.post("/api/convert", files=upload(flat_png(), "scan.png"))
        assert response.status_code == 200, response.text
        body = response.json()

        skips = [line for line in body["warnings"] if "NOT rendered" in line]
        assert len(skips) == 1, body["warnings"]
        # The figure quoted is the UNION of the two crops, which is what the
        # slide really would have carried.
        assert "90% of the slide" in skips[0], skips[0]
        # Nothing may claim the file is unaffected - this is the case where the
        # file was the broken part.
        assert not any("unaffected" in line for line in body["warnings"]), body[
            "warnings"
        ]

        deck = client.get(f"/api/download/{body['download_id']}")
        assert deck.status_code == 200
        presentation = assert_native_deck(deck.content)
        pictures = [
            shape
            for shape in presentation.slides[0].shapes
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        ]
        assert len(pictures) == 1
        assert pictures[0].image.blob, "the surviving crop's bytes are gone"

        with zipfile.ZipFile(io.BytesIO(deck.content)) as archive:
            media = [n for n in archive.namelist() if n.startswith("ppt/media/")]
        assert len(media) == 1, media

    def test_the_preview_reports_both_new_kinds_in_z_order(self, body):
        kinds = [shape["kind"] for shape in body["preview"]["shapes"]]
        assert kinds == [
            "shape",
            "shape",
            "shape",
            "connector",
            "connector",
            "image",
            "text",
        ]

    def test_the_connector_payload_is_the_pinned_contract(self, body):
        connector = body["preview"]["shapes"][3]
        assert connector == {
            "kind": "connector",
            "start_pct": {"x": 12.5, "y": 30.0},
            "end_pct": {"x": 48.0, "y": 30.0},
            "elbow": False,
            "arrow_start": False,
            "arrow_end": True,
            "dash": "solid",
            "color": "#1E3A6E",
            "width_pt": 1.5,
        }

    def test_the_elbow_connector_is_reported_as_dashed_and_bent(self, body):
        elbow = body["preview"]["shapes"][4]
        assert elbow["elbow"] is True
        assert elbow["dash"] == "dashed"
        assert elbow["start_pct"] == {"x": 60.0, "y": 10.0}
        assert elbow["end_pct"] == {"x": 90.0, "y": 60.0}

    def test_the_image_payload_carries_the_crop_and_its_geometry(self, body):
        image = body["preview"]["shapes"][5]
        assert image["kind"] == "image"
        assert image["left_pct"] == 60.0
        assert image["top_pct"] == 20.0
        assert image["width_pct"] == 30.0
        assert image["height_pct"] == 40.0
        assert image["data_uri"], "the conversion had the source image; it must crop"
        assert image["data_uri"].startswith("data:image/")
        assert len(image["data_uri"]) <= 96 * 1024

    def test_the_crop_came_from_the_right_region_of_the_upload(self, body):
        """The plumbing test with teeth: only the claimed region of the upload
        is red, so a crop that came from the wrong place - or from a blank
        image - is white, and a crop that never happened is ``null``."""
        from PIL import Image

        data_uri = body["preview"]["shapes"][5]["data_uri"]
        raw = base64.b64decode(data_uri.partition(",")[2])
        crop = Image.open(io.BytesIO(raw)).convert("RGB")

        for point in ((0, 0), (crop.width - 1, crop.height - 1)):
            for channel, expected in zip(crop.getpixel(point), IMAGE_RGB):
                assert abs(channel - expected) <= 8, (point, crop.getpixel(point))

    def test_nothing_was_dropped_and_the_counts_name_the_new_types(self, body):
        assert body["dropped"] == []
        assert body["counts"]["connector"] == 2
        assert body["counts"]["image"] == 1
        assert body["counts"]["shape"] == 3
        assert body["counts"]["text"] == 1
        assert body["total"] == 7

    def test_no_warning_claims_a_missing_source_image(self, body):
        """The honesty rule: the conversion HAD the image, so nothing may say
        it did not."""
        assert not [
            line for line in body["warnings"] if "no source image" in line
        ], body["warnings"]

    def test_the_spec_round_trips_with_both_new_element_types(self, body):
        """What the UI posts back to /api/preview has to validate again."""
        spec = SlideSpec.model_validate(body["spec"])
        assert spec.counts_by_type()["connector"] == 2
        assert spec.counts_by_type()["image"] == 1

    def test_a_dropped_crop_is_reported_in_the_responses_warnings(
        self, client, monkeypatch, with_key, no_client_allowed
    ):
        """``preview.py`` gives a crop up on stderr, like every other non-fatal
        problem in this project; ``warning_lines`` is what turns that into the
        response's ``warnings`` array. This asserts the whole plumbing, because
        a ``data_uri`` that silently went ``null`` would leave the panel with an
        unexplained placeholder - and the house rule is that the user is told.
        """
        monkeypatch.setattr(
            converter.preview_module, "PREVIEW_IMAGE_MAX_URI_BYTES", 64
        )
        body = diagram_preview(client, diagram_png(), monkeypatch)

        assert body["preview"]["shapes"][5]["data_uri"] is None
        lines = [line for line in body["warnings"] if "too large" in line]
        assert lines, body["warnings"]
        # The module prefix is stripped, like [pptx_builder] and [vision_analyzer].
        assert not any(line.startswith("[") for line in body["warnings"])
        # ...and the deck the user downloads still carries the real crop.
        deck = client.get(f"/api/download/{body['download_id']}")
        presentation = open_deck(deck.content)
        assert pptx_builder.flattening_report(presentation)["picture_count"] == 1


class TestConvertAndPreviewAgree:
    """One code path, asserted rather than assumed.

    Both endpoints build the deck with ``converter.build_deck`` and describe it
    with ``preview.build_preview``; the only difference is the source image. So
    for a spec with NO image element the two previews must be byte-identical -
    if they ever diverge, one of them grew its own derivation.
    """

    #: The diagram spec minus its image element: connectors and block arrows
    #: need no source image, so both endpoints can render all of it.
    @staticmethod
    def spec_without_images() -> dict:
        payload = copy.deepcopy(DIAGRAM_SPEC)
        payload["elements"] = [
            element for element in payload["elements"] if element["type"] != "image"
        ]
        return payload

    def test_the_two_previews_are_identical(
        self, client, monkeypatch, with_key, no_client_allowed
    ):
        payload = self.spec_without_images()
        monkeypatch.setattr(
            vision_analyzer,
            "analyze_image",
            fake_analyzer(SlideSpec.model_validate(copy.deepcopy(payload))),
        )
        converted = client.post(
            "/api/convert", files=upload(diagram_png(), "diagram.png")
        ).json()
        assert converted["preview"] is not None

        # Exactly what the UI does: post the spec it was handed back.
        previewed = client.post("/api/preview", json=converted["spec"]).json()

        assert previewed["preview"] == converted["preview"]
        assert previewed["counts"] == converted["counts"]
        assert previewed["warnings"] == converted["warnings"] == []

    def test_a_spec_with_an_image_diverges_only_in_that_element(
        self, client, monkeypatch, with_key, no_client_allowed
    ):
        """With an image element the two CANNOT agree - /api/preview has no
        source image - and the difference has to be confined to that element
        plus its warning, not smeared across the rest of the slide."""
        converted = diagram_preview(client, diagram_png(), monkeypatch)
        previewed = client.post("/api/preview", json=converted["spec"]).json()

        converted_kinds = [s["kind"] for s in converted["preview"]["shapes"]]
        previewed_kinds = [s["kind"] for s in previewed["preview"]["shapes"]]
        assert converted_kinds[5] == "image"
        assert previewed_kinds[5] == "shape", "the placeholder outline"
        assert converted_kinds[:5] == previewed_kinds[:5]
        assert converted_kinds[6:] == previewed_kinds[6:]

        # Every other entry is identical, element for element.
        for index in [0, 1, 2, 3, 4, 6]:
            assert (
                converted["preview"]["shapes"][index]
                == previewed["preview"]["shapes"][index]
            ), index

        assert converted["warnings"] == []
        assert any("no source image" in line for line in previewed["warnings"])


# =========================================================================== #
# POST /api/convert - every error code
# =========================================================================== #
class TestConvertErrors:
    @pytest.mark.parametrize("name", ["x.gif", "x.bmp", "x.tiff", "x.PNG.txt", "noext"])
    def test_unsupported_extension(self, client, png_bytes, name):
        response = client.post(
            "/api/convert", files=upload(png_bytes, name=name)
        )
        assert response.status_code == 400
        detail = envelope(response)
        assert detail["code"] == "unsupported_extension"
        for extension in sorted(config.SUPPORTED_EXTENSIONS):
            assert extension in detail["message"]

    def test_the_uploads_extension_decides_not_its_content_type(self, client, png_bytes):
        """A client can claim any content-type; the name is what is judged."""
        response = client.post(
            "/api/convert",
            files=upload(png_bytes, name="x.gif", content_type="image/png"),
        )
        assert response.status_code == 400
        assert envelope(response)["code"] == "unsupported_extension"

    def test_a_png_declared_as_octet_stream_is_still_accepted(
        self, client, png_bytes, offline_convert
    ):
        response = client.post(
            "/api/convert",
            files=upload(png_bytes, name="x.png", content_type="application/octet-stream"),
        )
        assert response.status_code == 200

    def test_empty_file(self, client):
        response = client.post(
            "/api/convert", files=upload(b"", name="empty.png")
        )
        assert response.status_code == 400
        assert envelope(response)["code"] == "empty_file"

    @pytest.mark.parametrize(
        "payload",
        [
            b"not an image, just text",
            b"<html><body>404 Not Found</body></html>",
            b"\x89PNG\r\n\x1a\n" + b"\x00" * 64,  # right magic, truncated body
        ],
    )
    def test_not_an_image(self, client, payload):
        response = client.post(
            "/api/convert", files=upload(payload, name="fake.png")
        )
        assert response.status_code == 400
        detail = envelope(response)
        assert detail["code"] == "not_an_image"
        # Pillow names the stream in its message; it must not be a heap address.
        assert "_io.BytesIO" not in detail["message"]
        assert "0x" not in detail["message"]

    def test_file_too_large_uses_the_real_20mb_cap(self, client):
        assert converter.MAX_UPLOAD_BYTES == 20 * 1024 * 1024
        oversize = b"\x89PNG\r\n\x1a\n" + b"\x00" * converter.MAX_UPLOAD_BYTES
        response = client.post(
            "/api/convert", files=upload(oversize, name="big.png")
        )
        assert response.status_code == 413
        detail = envelope(response)
        assert detail["code"] == "file_too_large"
        assert "larger than the 20 MB limit" in detail["message"]

    def test_the_cap_is_checked_before_the_image_is_decoded(self, client, monkeypatch):
        """Size beats decodability: junk bytes over the cap are a 413, not a 400,
        so a huge body is refused without being parsed as an image."""
        monkeypatch.setattr(converter, "MAX_UPLOAD_BYTES", 1024)
        response = client.post(
            "/api/convert",
            files=upload(b"z" * 4096, name="big.png"),
        )
        assert response.status_code == 413
        assert envelope(response)["code"] == "file_too_large"

    def test_a_file_exactly_at_the_cap_is_not_rejected_for_size(self, client, monkeypatch):
        monkeypatch.setattr(converter, "MAX_UPLOAD_BYTES", 1024)
        response = client.post(
            "/api/convert",
            files=upload(b"z" * 1024, name="edge.png"),
        )
        assert response.status_code == 400
        assert envelope(response)["code"] == "not_an_image"

    def test_missing_api_key(self, client, png_bytes):
        response = client.post("/api/convert", files=upload(png_bytes))
        assert response.status_code == 503
        detail = envelope(response)
        assert detail["code"] == "missing_api_key"
        assert "LITELLM_API_KEY" in detail["hint"]
        assert "export LITELLM_API_KEY" in detail["hint"]
        # The hint points at the two endpoints that genuinely need no key. It
        # must NOT point at mock mode, which no longer exists.
        assert "/api/build-from-json" in detail["hint"]
        assert "/api/preview" in detail["hint"]
        assert "mock" not in detail["hint"].lower()

    @pytest.mark.parametrize("blank", ["", "   ", "\t\n"])
    def test_a_blank_key_counts_as_missing(self, client, png_bytes, monkeypatch, blank):
        monkeypatch.setattr(config, "LITELLM_API_KEY", blank, raising=False)
        response = client.post("/api/convert", files=upload(png_bytes))
        assert response.status_code == 503
        assert envelope(response)["code"] == "missing_api_key"

    def test_the_key_check_happens_before_the_image_is_encoded(
        self, client, png_bytes, monkeypatch
    ):
        def _bomb(*args, **kwargs):  # pragma: no cover - a call is the failure
            raise AssertionError("must not analyze without a key")

        monkeypatch.setattr(vision_analyzer, "analyze_image", _bomb)
        assert client.post("/api/convert", files=upload(png_bytes)).status_code == 503

    def test_proxy_unreachable(self, client, png_bytes, with_key, monkeypatch):
        import httpx
        import openai

        def _unreachable(*args, **kwargs):
            try:
                raise openai.APIConnectionError(
                    request=httpx.Request("POST", "http://localhost:4000/")
                )
            except openai.APIConnectionError as exc:
                raise vision_analyzer.VisionAnalysisError("request failed") from exc

        monkeypatch.setattr(vision_analyzer, "analyze_image", _unreachable)
        response = client.post("/api/convert", files=upload(png_bytes))
        assert response.status_code == 502
        detail = envelope(response)
        assert detail["code"] == "proxy_unreachable"
        assert "LITELLM_BASE_URL" in detail["hint"]

    def test_analysis_failed(self, client, png_bytes, with_key, monkeypatch):
        def _bad_reply(*args, **kwargs):
            raise vision_analyzer.VisionAnalysisError(
                "Model 'gemini-flash' returned invalid JSON"
            )

        monkeypatch.setattr(vision_analyzer, "analyze_image", _bad_reply)
        response = client.post("/api/convert", files=upload(png_bytes))
        assert response.status_code == 502
        detail = envelope(response)
        assert detail["code"] == "analysis_failed"
        assert "returned invalid JSON" in detail["hint"]

    def test_a_missing_key_inside_the_analyzer_is_still_a_503(
        self, client, png_bytes, with_key, monkeypatch
    ):
        """The key is checked up front, but if the core raises it anyway (a race,
        or a client built elsewhere) the TYPE still decides the status code."""

        def _no_key(*args, **kwargs):
            try:
                raise config.MissingApiKeyError("no key")
            except config.MissingApiKeyError as exc:
                raise vision_analyzer.VisionAnalysisError("no key configured") from exc

        monkeypatch.setattr(vision_analyzer, "analyze_image", _no_key)
        response = client.post("/api/convert", files=upload(png_bytes))
        assert response.status_code == 503
        assert envelope(response)["code"] == "missing_api_key"

    def test_build_failed(self, client, png_bytes, offline_convert, monkeypatch):
        # ``**kwargs`` because the pipeline hands the builder the source image
        # as well now; a stand-in that refused it would fail with a TypeError
        # of its own and stop testing what it is named for.
        def _explode(spec, **kwargs):
            raise pptx_builder.PptxBuildError("python-pptx said no")

        monkeypatch.setattr(pptx_builder, "build_presentation", _explode)
        response = client.post("/api/convert", files=upload(png_bytes))
        assert response.status_code == 500
        detail = envelope(response)
        assert detail["code"] == "build_failed"
        assert "python-pptx said no" in detail["message"]

    def test_a_reply_with_no_usable_element_is_analysis_failed(
        self, client, png_bytes, with_key, monkeypatch
    ):
        """An all-dropped payload is a loud failure, never a blank deck - and on
        the convert path the blame is the model's, so it is a 502."""

        def _all_dropped(*args, **kwargs):
            SlideSpec.model_validate({"elements": [{"type": "table", "bbox": [0, 0, 1, 1]}]})

        monkeypatch.setattr(vision_analyzer, "analyze_image", _all_dropped)
        response = client.post("/api/convert", files=upload(png_bytes))
        assert response.status_code == 502
        detail = envelope(response)
        assert detail["code"] == "analysis_failed"
        assert "data - Field required" in detail["hint"]

    def test_a_missing_file_part_uses_the_same_envelope(self, client):
        response = client.post("/api/convert")
        assert response.status_code == 422
        assert envelope(response)["code"] == "invalid_request"

    def test_an_unexpected_crash_is_a_500_with_no_traceback(
        self, client, png_bytes, monkeypatch
    ):
        def _boom(*args, **kwargs):
            raise RuntimeError("secret internal detail at /srv/app/thing.py:42")

        monkeypatch.setattr(converter, "convert_upload", _boom)
        response = client.post("/api/convert", files=upload(png_bytes))
        assert response.status_code == 500
        detail = envelope(response)
        assert detail["code"] == "internal_error"
        assert "RuntimeError" in detail["message"]
        assert "secret internal detail" not in response.text
        assert "Traceback" not in response.text
        assert ".py" not in response.text


# =========================================================================== #
# POST /api/refine-instructions
# =========================================================================== #
#: A plausible completed instruction, in the note's language (Korean).
REFINED = (
    "우측 중앙의 3열 3행 표를 하나의 표로 인식하고, 첫 행을 헤더로 지정하세요."
)


class _FakeCompletions:
    """Records the call, so a test can assert what the route actually sent."""

    def __init__(self, reply):
        self.reply = reply
        self.calls: list[dict] = []

    def create(self, **kwargs):
        self.calls.append(kwargs)
        if isinstance(self.reply, BaseException):
            raise self.reply
        message = type("_M", (), {"content": self.reply, "refusal": None})()
        choice = type("_C", (), {"message": message, "finish_reason": "stop"})()
        return type("_R", (), {"choices": [choice]})()


class _FakeClient:
    def __init__(self, reply=REFINED):
        self.chat = type("_Chat", (), {})()
        self.chat.completions = _FakeCompletions(reply)

    @property
    def calls(self) -> list[dict]:
        return self.chat.completions.calls


@pytest.fixture
def offline_refine(monkeypatch, with_key):
    """The real ``/api/refine-instructions`` path, offline.

    ``config.get_client`` is what ``instructions.refine_instruction`` calls
    when no client is injected, so replacing it exercises the whole route -
    form parsing, upload validation, context selection and the response
    envelope - without the network. Returns the fake, so a test can assert on
    the prompt that was built.
    """
    fake = _FakeClient()
    monkeypatch.setattr(config, "get_client", lambda: fake)
    return fake


class TestRefineInstructions:
    """The HTTP layer of the instruction assist.

    ``tests/test_instructions.py`` covers the service in depth and is
    framework-free by design, so it cannot reach any of this: the multipart
    parsing, the optional file part, the order the route validates things in,
    or the status codes. Those live here.
    """

    URL = "/api/refine-instructions"

    def test_a_note_and_a_spec_answer_the_pinned_contract(
        self, client, offline_refine
    ):
        response = client.post(
            self.URL,
            data={"note": "표가 이상해", "spec": json.dumps(spec_dict())},
        )

        assert response.status_code == 200
        body = response.json()
        assert set(body) == {
            "status",
            "instruction",
            "note",
            "context_used",
            # Beside ``context_used``, not inside it: the three pinned context
            # values keep their meaning, and this says whether the model
            # produced anything at all. Without it the UI cannot tell a genuine
            # no-context run from an echo of the user's own note, and it prints
            # a sentence about each.
            "fell_back",
        }
        assert body["status"] == "success"
        assert body["instruction"] == REFINED
        # Echoed EXACTLY as received - the UI needs it to offer an undo.
        assert body["note"] == "표가 이상해"
        assert body["context_used"] == "spec"
        assert body["fell_back"] is False

    def test_the_spec_is_summarized_not_dumped(self, client, offline_refine):
        """The whole point of the digest: a spec goes up as a short
        description, so the prompt must not contain the bbox integers that the
        raw JSON is mostly made of."""
        spec = spec_dict()
        response = client.post(
            self.URL, data={"note": "표가 이상해", "spec": json.dumps(spec)}
        )
        assert response.status_code == 200

        sent = str(offline_refine.calls[0]["messages"][1]["content"])
        assert '"bbox"' not in sent
        assert "coord" not in sent.lower()
        # ...and it really did carry the slide, not just the note.
        assert "element" in sent.lower()

    def test_an_image_is_used_when_there_is_no_spec(
        self, client, png_bytes, offline_refine
    ):
        response = client.post(
            self.URL, data={"note": "표가 이상해"}, files=upload(png_bytes)
        )

        assert response.status_code == 200
        assert response.json()["context_used"] == "image"
        # One vision call, carrying the image as a data URL part.
        content = offline_refine.calls[0]["messages"][1]["content"]
        assert isinstance(content, list)
        assert any(part.get("type") == "image_url" for part in content)

    def test_a_bare_note_still_answers(self, client, offline_refine):
        response = client.post(self.URL, data={"note": "표가 이상해"})

        assert response.status_code == 200
        body = response.json()
        assert body["context_used"] == "none"
        assert body["instruction"]

    @pytest.mark.parametrize(
        "note",
        ["", "   ", "\t\n  \n", "　"],  # incl. an ideographic space
        ids=["empty", "spaces", "newlines", "ideographic-space"],
    )
    def test_a_blank_note_is_empty_note(self, client, offline_refine, note):
        response = client.post(self.URL, data={"note": note})

        assert response.status_code == 400
        assert envelope(response)["code"] == "empty_note"
        # Nothing was sent upstream for a request that cannot be answered.
        assert offline_refine.calls == []

    def test_a_missing_note_field_is_also_empty_note(self, client, offline_refine):
        """Not fastapi's own 422: the contract pins ONE code for "no note",
        so the UI has one Korean sentence for it."""
        response = client.post(self.URL, data={"spec": "{}"})

        assert response.status_code == 400
        assert envelope(response)["code"] == "empty_note"

    @pytest.mark.parametrize(
        "spec",
        [
            "}{ not json {{",
            '{"elements":[{"type":"text"',  # truncated mid-object
            "12345",
            "null",
            '{"elements":[]}',
            '{"elements":"not a list"}',
            "[]",
        ],
        ids=[
            "garbage",
            "truncated",
            "bare-number",
            "null",
            "no-elements",
            "elements-not-a-list",
            "empty-array",
        ],
    )
    def test_an_unusable_spec_is_ignored_not_rejected(
        self, client, offline_refine, spec
    ):
        """A spec is only CONTEXT. An unreadable one degrades to the next
        context down and SAYS so - it must never 500, and must never cost the
        user their answer."""
        response = client.post(self.URL, data={"note": "표가 이상해", "spec": spec})

        assert response.status_code == 200, response.text
        body = response.json()
        assert body["context_used"] == "none"
        assert body["instruction"]

    def test_an_unusable_spec_falls_back_to_the_image(
        self, client, png_bytes, offline_refine
    ):
        """Cheapest-first, but not at the cost of context: when the spec
        cannot be read and an image was sent, the image is used."""
        response = client.post(
            self.URL,
            data={"note": "표가 이상해", "spec": "}{ not json {{"},
            files=upload(png_bytes),
        )

        assert response.status_code == 200
        assert response.json()["context_used"] == "image"

    def test_a_non_image_file_is_not_an_image(self, client, offline_refine):
        """The optional context file goes through the same validation
        ``/api/convert`` uses - it is not re-implemented, so it cannot drift."""
        response = client.post(
            self.URL,
            data={"note": "표가 이상해"},
            files={"file": ("slide.png", b"this is plain text", "image/png")},
        )

        assert response.status_code == 400
        assert envelope(response)["code"] == "not_an_image"

    def test_an_oversized_file_reuses_the_upload_cap(self, client, offline_refine):
        oversized = b"A" * (converter.MAX_UPLOAD_BYTES + 1)
        response = client.post(
            self.URL,
            data={"note": "표가 이상해"},
            files={"file": ("slide.png", oversized, "image/png")},
        )

        assert response.status_code == 413
        assert envelope(response)["code"] == "file_too_large"

    def test_a_blank_note_beats_a_bad_file(self, client, offline_refine):
        """Two things are wrong; the code names the one the caller actually
        forgot. The note is required, the file is optional context."""
        response = client.post(
            self.URL,
            data={"note": "   "},
            files={"file": ("slide.png", b"not an image", "image/png")},
        )

        assert response.status_code == 400
        assert envelope(response)["code"] == "empty_note"

    def test_an_empty_file_part_means_no_file(self, client, offline_refine):
        """A browser posting a form with an untouched ``<input type=file>``
        still sends the part, with no bytes and no filename. That is "I have no
        image yet", not ``empty_file`` - answering an error there would break
        the assist for every user who has not picked a file."""
        response = client.post(
            self.URL, data={"note": "표가 이상해"}, files={"file": ("", b"", "")}
        )

        assert response.status_code == 200
        assert response.json()["context_used"] == "none"

    def test_a_named_but_empty_file_is_still_empty_file(self, client, offline_refine):
        """The other side of that boundary. A part with a FILENAME is someone
        genuinely trying to send an image, so zero bytes is a real problem
        worth naming - the "no file at all" path above is the one that must
        stay silent."""
        response = client.post(
            self.URL,
            data={"note": "표가 이상해"},
            files={"file": ("slide.png", b"", "image/png")},
        )

        assert response.status_code == 400
        assert envelope(response)["code"] == "empty_file"

    def test_no_key_is_missing_api_key(self, client, monkeypatch):
        """Checked before anything is encoded or sent: 503, by name, so the
        operator learns which variable to set."""
        monkeypatch.setattr(config, "LITELLM_API_KEY", "", raising=False)

        def _bomb():  # pragma: no cover - a call is the failure
            raise AssertionError("no client may be built without a key")

        monkeypatch.setattr(config, "get_client", _bomb)

        response = client.post(self.URL, data={"note": "표가 이상해"})

        assert response.status_code == 503
        assert envelope(response)["code"] == "missing_api_key"

    def test_an_upstream_failure_is_refine_failed(self, client, monkeypatch, with_key):
        fake = _FakeClient(reply=RuntimeError("upstream exploded"))
        monkeypatch.setattr(config, "get_client", lambda: fake)

        response = client.post(self.URL, data={"note": "표가 이상해"})

        assert response.status_code == 502
        assert envelope(response)["code"] == "refine_failed"

    def test_build_failed_is_not_reachable_here(self, client, monkeypatch, with_key):
        """Nothing is built on this path, so the contract deliberately leaves
        ``build_failed`` out of it."""
        fake = _FakeClient(reply=RuntimeError("upstream exploded"))
        monkeypatch.setattr(config, "get_client", lambda: fake)

        response = client.post(self.URL, data={"note": "표가 이상해"})
        assert envelope(response)["code"] != "build_failed"

    @pytest.mark.parametrize(
        "reply",
        ["", "   ", "```\n```", "...", None],
        ids=["empty", "spaces", "empty-fence", "punctuation", "none"],
    )
    def test_a_useless_reply_never_empties_the_textarea(
        self, client, monkeypatch, with_key, reply
    ):
        """The one thing this feature must not do. The frontend writes
        ``instruction`` straight into the box the user typed in, so a 200 with
        an empty instruction would delete their own words."""
        fake = _FakeClient(reply=reply)
        monkeypatch.setattr(config, "get_client", lambda: fake)

        response = client.post(self.URL, data={"note": "표가 이상해"})

        assert response.status_code == 200
        body = response.json()
        assert body["instruction"].strip()
        # It is the note itself coming back, so the honest answer is "none":
        # neither spec nor image informed it.
        assert body["context_used"] == "none"

    def test_a_fallback_does_not_claim_the_spec_was_used(
        self, client, monkeypatch, with_key
    ):
        """A usable spec WAS sent, but the model gave nothing back, so the
        answer is only the tidied note. Reporting ``spec`` here would make the
        UI's "슬라이드 분석 결과를 참고했습니다" a lie."""
        fake = _FakeClient(reply="")
        monkeypatch.setattr(config, "get_client", lambda: fake)

        response = client.post(
            self.URL,
            data={"note": "표가 이상해", "spec": json.dumps(spec_dict())},
        )

        assert response.status_code == 200
        assert response.json()["context_used"] == "none"

    def test_the_model_field_defaults_and_overrides(self, client, offline_refine):
        """Same as ``/api/convert``: absent means the server's configured
        model, present means exactly what was asked for."""
        client.post(self.URL, data={"note": "표가 이상해"})
        assert offline_refine.calls[0]["model"] == config.DEFAULT_MODEL

        client.post(self.URL, data={"note": "표가 이상해", "model": "other-model"})
        assert offline_refine.calls[1]["model"] == "other-model"

    def test_the_endpoint_is_in_the_schema_with_its_fields(self, client):
        schema = client.get("/openapi.json").json()
        post = schema["paths"]["/api/refine-instructions"]["post"]
        ref = post["requestBody"]["content"]["multipart/form-data"]["schema"]["$ref"]
        fields = schema["components"]["schemas"][ref.rsplit("/", 1)[-1]]["properties"]

        assert set(fields) >= {"note", "spec", "file", "model"}

    def test_the_instruction_is_length_capped(self, client, monkeypatch, with_key):
        """An ``extra_instructions`` value longer than the analyzer's own
        prompt stops being guidance, and it has to fit a textarea a human can
        still read."""
        fake = _FakeClient(reply="가" * 5000)
        monkeypatch.setattr(config, "get_client", lambda: fake)

        response = client.post(self.URL, data={"note": "표가 이상해"})

        assert response.status_code == 200
        assert len(response.json()["instruction"]) <= 600


# =========================================================================== #
# POST /api/build-from-json
# =========================================================================== #
class TestBuildFromJson:
    def post(self, client, payload, **params):
        return client.post(
            "/api/build-from-json",
            content=json.dumps(payload),
            headers={"content-type": "application/json"},
            params=params or None,
        )

    def test_a_plain_slidespec(self, client):
        response = self.post(client, spec_dict())
        assert response.status_code == 200
        assert response.headers["content-type"] == PPTX_MEDIA_TYPE
        assert response.headers["content-disposition"] == 'attachment; filename="slide.pptx"'
        assert json.loads(response.headers["x-element-counts"]) == {
            "shape": 4,
            "table": 1,
            "text": 4,
        }
        assert response.headers["x-dropped-count"] == "0"
        assert_native_deck(response.content)

    def test_a_spec_under_a_wrapper_key(self, client):
        response = self.post(client, {"spec": spec_dict()})
        assert response.status_code == 200
        assert json.loads(response.headers["x-element-counts"])["shape"] == 4
        assert_native_deck(response.content)

    def test_a_bare_element_array(self, client):
        elements = spec_dict()["elements"]
        response = self.post(client, elements)
        assert response.status_code == 200
        assert json.loads(response.headers["x-element-counts"]) == {
            "shape": 4,
            "table": 1,
            "text": 4,
        }
        assert_native_deck(response.content)

    def test_a_partial_drop_is_a_200_that_reports_the_loss(self, client):
        payload = spec_dict()
        payload["elements"].append({"type": "table", "bbox": [10, 10, 100, 100]})
        payload["elements"].append({"type": "text", "bbox": [0, 0, 10, 10]})  # no content

        response = self.post(client, payload)
        assert response.status_code == 200
        assert response.headers["x-dropped-count"] == "2"
        assert json.loads(response.headers["x-element-counts"]) == {
            "shape": 4,
            "table": 1,
            "text": 4,
        }
        assert_native_deck(response.content)

    def test_an_incoming_dropped_key_cannot_forge_diagnostics(self, client):
        payload = spec_dict()
        payload["dropped"] = ["element 99 (text): invented loss"]
        response = self.post(client, payload)
        assert response.status_code == 200
        assert response.headers["x-dropped-count"] == "0"

    def test_the_filename_query_param(self, client):
        response = self.post(client, spec_dict(), filename="Quarterly Review")
        assert response.status_code == 200
        assert (
            response.headers["content-disposition"]
            == 'attachment; filename="Quarterly Review.pptx"'
        )

    def test_a_hostile_filename_cannot_inject_a_header(self, client):
        response = self.post(client, spec_dict(), filename='../../x"; evil="1')
        assert response.status_code == 200
        disposition = response.headers["content-disposition"]
        assert disposition == 'attachment; filename="x_evil_1.pptx"'
        assert disposition.count('"') == 2
        assert disposition.count(";") == 1
        assert "/" not in disposition.split("filename=")[1]

    def test_a_hangul_filename_query_param_survives_the_round_trip(self, client):
        """The Spec JSON tab's Download button sends the name it is showing. It
        used to be sanitized twice - once in the route, once in
        ``pptx_headers`` - so the non-ASCII was already gone before the header
        was built and every such download was called "slide.pptx"."""
        response = self.post(
            client, spec_dict(), filename="분기보고서.pptx"
        )
        assert response.status_code == 200
        disposition = response.headers["content-disposition"]
        assert disposition.startswith('attachment; filename="slide.pptx"; ')
        encoded = disposition.split("filename*=UTF-8''", 1)[1]
        assert (
            urllib.parse.unquote(encoded) == "분기보고서.pptx"
        )
        assert_native_deck(response.content)

    @pytest.mark.parametrize(
        "payload",
        [
            {"elements": [{"type": "table", "bbox": [0, 0, 100, 100]}]},
            [{"type": "text", "bbox": [0, 0, 100, 100]}],  # no content -> dropped
        ],
    )
    def test_invalid_spec_when_every_element_is_dropped(self, client, payload):
        response = self.post(client, payload)
        assert response.status_code == 400
        detail = envelope(response)
        assert detail["code"] == "invalid_spec"
        assert "element 0" in detail["hint"]

    def test_invalid_spec_lists_every_drop_reason_in_the_hint(self, client):
        response = self.post(
            client,
            [
                {"type": "table", "bbox": [0, 0, 100, 100]},
                {"type": "text", "bbox": [0, 0, 100, 100]},
            ],
        )
        assert response.status_code == 400
        hint = envelope(response)["hint"]
        assert "element 0" in hint and "element 1" in hint
        assert "2 element(s)" in hint

    def test_invalid_spec_for_a_wrong_shaped_body(self, client):
        response = self.post(client, {"elements": "not a list"})
        assert response.status_code == 400
        assert envelope(response)["code"] == "invalid_spec"

    def test_invalid_spec_for_a_body_that_is_not_json(self, client):
        response = client.post(
            "/api/build-from-json",
            content=b"{not json at all",
            headers={"content-type": "application/json"},
        )
        assert response.status_code == 400
        assert envelope(response)["code"] == "invalid_spec"

    def test_invalid_spec_for_an_empty_body(self, client):
        response = client.post(
            "/api/build-from-json",
            content=b"",
            headers={"content-type": "application/json"},
        )
        assert response.status_code == 400
        assert envelope(response)["code"] == "invalid_spec"

    def test_build_failed(self, client, monkeypatch):
        def _explode(spec, **kwargs):
            raise pptx_builder.PptxBuildError("cannot render")

        monkeypatch.setattr(pptx_builder, "build_presentation", _explode)
        response = self.post(client, spec_dict())
        assert response.status_code == 500
        assert envelope(response)["code"] == "build_failed"

    def test_an_empty_spec_builds_an_empty_deck(self, client):
        """No elements at all is not a drop, so it is a legitimate 200."""
        response = self.post(client, {"background_color": "#FFFFFF", "elements": []})
        assert response.status_code == 200
        assert response.headers["x-dropped-count"] == "0"
        assert json.loads(response.headers["x-element-counts"]) == {
            "shape": 0,
            "table": 0,
            "text": 0,
        }
        assert len(assert_native_deck(response.content).slides[0].shapes) == 0

    def test_it_does_not_register_a_download(self, client):
        self.post(client, spec_dict())
        assert converter.downloads.ids() == []


# =========================================================================== #
# GET /api/download/{file_id}
# =========================================================================== #
class TestDownload:
    def _convert(self, client, png_bytes):
        response = client.post("/api/convert", files=upload(png_bytes))
        assert response.status_code == 200
        return response.json()

    def test_happy_path_headers_and_body(self, client, png_bytes, offline_convert):
        body = self._convert(client, png_bytes)
        response = client.get(f"/api/download/{body['download_id']}")

        assert response.status_code == 200
        assert response.headers["content-type"] == PPTX_MEDIA_TYPE
        assert response.headers["content-disposition"] == 'attachment; filename="sample.pptx"'
        assert json.loads(response.headers["x-element-counts"]) == body["counts"]
        assert response.headers["x-dropped-count"] == str(len(body["dropped"]))
        assert_native_deck(response.content)

    def test_a_hangul_upload_name_reaches_the_download_header(
        self, client, png_bytes, offline_convert
    ):
        """End to end: the store has to keep the user's own characters, because
        the header is built from what it kept. Sanitizing on the way IN threw
        them away, so a Korean-named upload was called "slide.pptx" both in the
        response body the UI shows and in the saved file - for the very users
        this UI is in Korean for."""
        body = client.post(
            "/api/convert",
            files=upload(png_bytes, name="매출표_Q3.png"),
        ).json()

        disposition = client.get(f"/api/download/{body['download_id']}").headers[
            "content-disposition"
        ]
        # The body carries the name the user recognises (UTF-8 JSON), and the
        # header carries BOTH forms - the latin-1 fallback and the real one.
        assert body["filename"] == "매출표_Q3.pptx"
        assert disposition.startswith('attachment; filename="Q3.pptx"; ')
        encoded = disposition.split("filename*=UTF-8''", 1)[1]
        assert urllib.parse.unquote(encoded) == body["filename"]
        disposition.encode("latin-1")
        assert "\r" not in disposition and "\n" not in disposition

    def test_an_ascii_upload_name_still_gets_the_one_parameter_header(
        self, client, png_bytes, offline_convert
    ):
        body = client.post(
            "/api/convert", files=upload(png_bytes, name="review-slide.png")
        ).json()
        disposition = client.get(f"/api/download/{body['download_id']}").headers[
            "content-disposition"
        ]
        assert disposition == 'attachment; filename="review-slide.pptx"'
        assert body["filename"] == "review-slide.pptx"

    def test_re_downloading_the_same_id_works(self, client, png_bytes, offline_convert):
        body = self._convert(client, png_bytes)
        first = client.get(f"/api/download/{body['download_id']}")
        second = client.get(f"/api/download/{body['download_id']}")
        third = client.get(f"/api/download/{body['download_id']}")

        assert [first.status_code, second.status_code, third.status_code] == [200, 200, 200]
        assert first.content == second.content == third.content
        assert len(first.content) > 0

    def test_unknown_id(self, client):
        unknown = uuid.uuid4().hex
        response = client.get(f"/api/download/{unknown}")
        assert response.status_code == 404
        detail = envelope(response)
        assert detail["code"] == "unknown_download"
        assert "conversion" in detail["hint"]

    @pytest.mark.parametrize(
        "bad_id",
        [
            "not-a-uuid",
            "0" * 32,  # hex, right length, but the uuid version nibble is 0
            uuid.uuid1().hex,  # a real uuid, wrong version
            "ZZ" + uuid.uuid4().hex[2:],
            uuid.uuid4().hex.upper(),
            uuid.uuid4().hex[:31],
            uuid.uuid4().hex + "a",
            str(uuid.uuid4()),  # dashed form
            "....",
            "%2e%2e",
            "..etc..passwd..",
            "etc-passwd",
        ],
    )
    def test_a_non_uuid4_id_is_rejected(self, client, bad_id):
        response = client.get(f"/api/download/{bad_id}")
        assert response.status_code == 404
        assert envelope(response)["code"] == "unknown_download"

    def test_a_non_uuid_id_never_reaches_the_store_or_the_filesystem(
        self, client, monkeypatch
    ):
        """The id comes straight from the URL, so it must be validated BEFORE
        any lookup - and it must never be interpolated into a path."""
        import builtins

        def _no_lookup(file_id):  # pragma: no cover - a call is the failure
            raise AssertionError(f"store was queried with {file_id!r}")

        monkeypatch.setattr(converter.downloads, "get", _no_lookup)

        opened: list[str] = []
        real_open = builtins.open

        def _recording_open(file, *args, **kwargs):
            opened.append(str(file))
            return real_open(file, *args, **kwargs)

        monkeypatch.setattr(builtins, "open", _recording_open)

        # One URL segment (so it really reaches the route) that is full of the
        # characters a path-traversal attempt is made of.
        marker = "passwd-marker"
        response = client.get(f"/api/download/..{marker}..")
        assert response.status_code == 404
        assert envelope(response)["code"] == "unknown_download"
        assert not [path for path in opened if marker in path]

    def test_two_conversions_get_independent_ids(self, client, png_bytes, offline_convert):
        first = self._convert(client, png_bytes)
        second = self._convert(client, png_bytes)
        assert first["download_id"] != second["download_id"]
        assert client.get(f"/api/download/{first['download_id']}").status_code == 200
        assert client.get(f"/api/download/{second['download_id']}").status_code == 200

    def test_a_download_carries_its_own_conversions_filename(
        self, client, png_bytes, offline_convert
    ):
        first = client.post(
            "/api/convert", files=upload(png_bytes, name="alpha.png")
        ).json()
        second = client.post(
            "/api/convert", files=upload(png_bytes, name="beta.jpg")
        ).json()

        assert 'filename="alpha.pptx"' in client.get(
            f"/api/download/{first['download_id']}"
        ).headers["content-disposition"]
        assert 'filename="beta.pptx"' in client.get(
            f"/api/download/{second['download_id']}"
        ).headers["content-disposition"]


# =========================================================================== #
# the download store itself
# =========================================================================== #
class TestDownloadStore:
    def test_the_contract_defaults(self):
        assert converter.DOWNLOAD_TTL_SECONDS == 30 * 60
        assert converter.DOWNLOAD_MAX_ENTRIES == 64
        assert converter.downloads.ttl == 30 * 60
        assert converter.downloads.max_entries == 64

    def test_put_returns_a_uuid4_hex(self):
        store = DownloadStore()
        file_id = store.put(b"deck", "a.pptx")
        assert converter.is_download_id(file_id)
        assert uuid.UUID(hex=file_id).version == 4

    def test_get_does_not_consume(self):
        store = DownloadStore()
        file_id = store.put(b"deck", "a.pptx", counts={"text": 1}, dropped_count=2)
        for _ in range(5):
            entry = store.get(file_id)
            assert entry is not None
            assert entry.data == b"deck"
            assert entry.filename == "a.pptx"
            assert entry.counts == {"text": 1}
            assert entry.dropped_count == 2
        assert len(store) == 1

    def test_an_entry_expires_after_the_ttl(self):
        store = DownloadStore(ttl=0.02)
        file_id = store.put(b"deck", "a.pptx")
        assert store.get(file_id) is not None
        time.sleep(0.05)
        assert store.get(file_id) is None
        assert len(store) == 0

    def test_the_cap_evicts_the_oldest_first(self):
        store = DownloadStore(max_entries=3)
        ids = [store.put(bytes([index]), f"{index}.pptx") for index in range(5)]

        assert len(store) == 3
        assert store.ids() == ids[2:]
        assert store.get(ids[0]) is None
        assert store.get(ids[1]) is None
        for file_id in ids[2:]:
            assert store.get(file_id) is not None

    def test_it_never_grows_past_the_cap(self):
        store = DownloadStore(max_entries=4)
        for index in range(50):
            store.put(b"x", f"{index}.pptx")
            assert len(store) <= 4

    def test_get_refuses_a_malformed_id_without_a_lookup(self):
        store = DownloadStore()
        store.put(b"deck", "a.pptx")
        for bad in [None, 42, "", "../etc/passwd", "0" * 32, uuid.uuid1().hex]:
            assert store.get(bad) is None

    def test_clear_empties_it(self):
        store = DownloadStore()
        store.put(b"deck", "a.pptx")
        store.clear()
        assert len(store) == 0


# =========================================================================== #
# the converter seam
# =========================================================================== #
class TestConverterSeam:
    def test_it_imports_no_web_framework(self):
        """The seam must stay usable without fastapi, so no import line in it
        may mention a web framework (the prose in its docstrings may)."""
        import ast
        import inspect

        imported: set[str] = set()
        for node in ast.walk(ast.parse(inspect.getsource(converter))):
            if isinstance(node, ast.Import):
                imported.update(alias.name.split(".")[0] for alias in node.names)
            elif isinstance(node, ast.ImportFrom) and node.module and node.level == 0:
                imported.add(node.module.split(".")[0])

        assert imported, "the AST walk found no imports at all"
        assert not imported & {"fastapi", "starlette", "flask", "django"}

    def test_the_deck_is_built_in_memory_not_on_disk(self, monkeypatch):
        """``build_pptx`` writes to a path; the API must not use it."""

        def _bomb(*args, **kwargs):  # pragma: no cover - a call is the failure
            raise AssertionError("build_pptx() writes to disk; use build_presentation")

        monkeypatch.setattr(pptx_builder, "build_pptx", _bomb)
        data = converter.build_bytes(mock_data.sample_spec())
        assert data[:2] == b"PK"
        assert_native_deck(data)

    def test_build_bytes_is_deterministic_enough_to_reopen(self):
        data = converter.build_bytes(mock_data.sample_spec())
        assert len(open_deck(data).slides) == 1

    def test_counts_always_carry_the_three_known_types(self):
        spec = SlideSpec.model_validate(
            {"elements": [{"type": "text", "bbox": [0, 0, 10, 10], "content": "hi"}]}
        )
        assert converter.counts_for(spec) == {"shape": 0, "table": 0, "text": 1}

    def test_counts_match_the_cores_own_tally(self):
        spec = mock_data.sample_spec()
        core = spec.counts_by_type()
        counts = converter.counts_for(spec)
        for name, value in core.items():
            assert counts[name] == value
        assert sum(counts.values()) == len(spec.elements)

    # ----------------------------------------------------------------- #
    # the non-ASCII half of the download name
    # ----------------------------------------------------------------- #
    @pytest.mark.parametrize(
        ("raw", "expected"),
        [
            # What safe_pptx_filename destroys and this has to keep.
            ("분기버전.png", "분기버전.pptx"),
            ("매출표_Q3.png", "매출표_Q3.pptx"),
            ("2026 분기버전.png", "2026 분기버전.pptx"),
            ("éèê.png", "éèê.pptx"),
            # ...while the ASCII cases stay exactly what the header already said.
            ("deck.png", "deck.pptx"),
            ("deck.pptx", "deck.pptx"),
            ("a/b/c/deck.jpeg", "deck.pptx"),
            ("a\\b\\deck.webp", "deck.pptx"),
            ("", "slide.pptx"),
            (None, "slide.pptx"),
            ("...", "slide.pptx"),
        ],
    )
    def test_display_pptx_filename_keeps_the_users_own_letters(self, raw, expected):
        assert converter.display_pptx_filename(raw) == expected

    @pytest.mark.parametrize(
        "raw",
        [
            'ev"il\r\nX; q=1.png',
            "분기\r\n버전.png",
            "a\x00b/../c.png",
            "매출\u007f\u0085표.png",
        ],
    )
    def test_the_display_name_still_carries_nothing_structural(self, raw):
        """It is percent-encoded into the header, so a CRLF could not survive
        anyway - but it also ends up as a filename on the client."""
        name = converter.display_pptx_filename(raw)
        assert not set(name) & set('"\\/:;*?<>|\r\n\x00')
        assert name.endswith(".pptx")

    def test_a_hangul_upload_gets_both_disposition_parameters(self):
        """``filename=`` is latin-1 only, so a Hangul name had nowhere to go and
        every such download was called "slide.pptx". RFC 5987's ``filename*``
        is the parameter that carries it, and browsers prefer it."""
        header = converter.pptx_headers(
            "분기보고서.png", {"text": 1}, 0
        )["Content-Disposition"]

        # The latin-1 fallback is unchanged and still safe to put in a header.
        assert header.startswith('attachment; filename="slide.pptx"; ')
        header.encode("latin-1")
        assert "\r" not in header and "\n" not in header
        assert header.count('"') == 2

        encoded = header.split("filename*=UTF-8''", 1)[1]
        assert urllib.parse.unquote(encoded) == "분기보고서.pptx"

    def test_an_ascii_upload_gets_exactly_the_header_it_always_did(self):
        """``filename*`` is emitted only when it adds information, so the common
        case is byte-identical - including for a hostile name, whose every
        offending character is ASCII."""
        for raw in ("sample.png", "my slide.jpg", 'ev"il\r\nX; q=1.png'):
            header = converter.pptx_headers(raw, {}, 0)["Content-Disposition"]
            assert "filename*" not in header
            assert header.count(";") == 1
            assert header == f'attachment; filename="{converter.safe_pptx_filename(raw)}"'

    def test_a_hostile_hangul_name_cannot_inject_a_header(self):
        """The one case where both parameters AND hostile input meet."""
        header = converter.pptx_headers(
            '분기"; x=1\r\nX-Evil: yes.png', {}, 0
        )["Content-Disposition"]
        header.encode("latin-1")
        assert "\r" not in header and "\n" not in header
        assert header.count('"') == 2
        # The percent-encoded half contains no raw separator either.
        encoded = header.split("filename*=UTF-8''", 1)[1]
        assert not set(encoded) & set('";, \r\n')

    @pytest.mark.parametrize(
        ("raw", "expected"),
        [
            ("deck.png", "deck.pptx"),
            ("deck.pptx", "deck.pptx"),
            ("a/b/c/deck.jpeg", "deck.pptx"),
            ("a\\b\\deck.webp", "deck.pptx"),
            ("", "slide.pptx"),
            (None, "slide.pptx"),
            ("...", "slide.pptx"),
            ("\u00e9\u00e8\u00ea.png", "slide.pptx"),
            ("a" * 300 + ".png", "a" * 100 + ".pptx"),
        ],
    )
    def test_safe_pptx_filename(self, raw, expected):
        name = converter.safe_pptx_filename(raw)
        assert name == expected
        name.encode("latin-1")  # a Content-Disposition value must survive this

    def test_failure_from_exception_branches_on_type_not_message(self):
        """A reworded core message must not be able to change a status code."""
        cases = [
            (config.MissingApiKeyError("anything at all"), 503, "missing_api_key"),
            (pptx_builder.PptxBuildError("anything at all"), 500, "build_failed"),
            (vision_analyzer.VisionAnalysisError("anything at all"), 502, "analysis_failed"),
        ]
        for exc, status, code in cases:
            failure = converter.failure_from_exception(exc, context="convert")
            assert (failure.status_code, failure.code) == (status, code)

    @pytest.mark.parametrize("context", ["convert", "build"])
    def test_an_unknown_exception_stays_inside_the_contracts_error_table(
        self, context
    ):
        """The error map may not invent a code a client was never told about.

        Both endpoints pin exactly one 500 - ``build_failed`` - so an exception
        type nobody anticipated has to be reported with that code rather than
        with the off-contract ``internal_error`` this branch used to return on
        the convert path. ``internal_error`` stays reachable only from
        ``app.main``'s handler for a crash OUTSIDE the mapped pipeline.
        """
        failure = converter.failure_from_exception(
            KeyError("secret internal detail"), context=context
        )
        assert (failure.status_code, failure.code) == (500, "build_failed")
        assert "KeyError" in failure.message
        # The class name is diagnostic; the message text could be anything.
        assert "secret internal detail" not in failure.hint

    def test_the_same_bad_spec_is_400_on_build_and_502_on_convert(self):
        try:
            SlideSpec.model_validate({"elements": [{"type": "table", "bbox": [0, 0, 1, 1]}]})
        except Exception as exc:  # noqa: BLE001
            build = converter.failure_from_exception(exc, context="build")
            convert = converter.failure_from_exception(exc, context="convert")
        assert (build.status_code, build.code) == (400, "invalid_spec")
        assert (convert.status_code, convert.code) == (502, "analysis_failed")
        assert "data - Field required" in build.hint

    def test_a_conversion_failure_passes_through_unchanged(self):
        original = ConversionFailure(418, "teapot", "short and stout", "tip me over")
        assert converter.failure_from_exception(original) is original
        assert original.detail == {
            "code": "teapot",
            "message": "short and stout",
            "hint": "tip me over",
        }

    def test_capture_stderr_is_scoped_to_the_calling_thread(self):
        """Two conversions run side by side in fastapi's threadpool; neither may
        see the other's warnings."""
        import sys
        import threading

        seen: dict[str, str] = {}
        started = threading.Barrier(2)

        def worker(name: str) -> None:
            with converter.capture_stderr() as buffer:
                started.wait(timeout=5)
                print(f"[pptx_builder] warning: from {name}", file=sys.stderr)
                time.sleep(0.01)
                seen[name] = buffer.getvalue()

        threads = [threading.Thread(target=worker, args=(name,)) for name in ("a", "b")]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(timeout=5)

        assert seen["a"].strip() == "[pptx_builder] warning: from a"
        assert seen["b"].strip() == "[pptx_builder] warning: from b"

    def test_capture_stderr_restores_the_original_stream(self):
        import sys

        before = sys.stderr
        with converter.capture_stderr():
            assert sys.stderr is not before
        assert sys.stderr is before

    def test_warning_lines_strips_prefixes_and_blank_lines(self):
        text = (
            "[pptx_builder] warning: first thing\n"
            "\n"
            "[vision_analyzer] second thing\n"
            "   \n"
            "bare line\n"
        )
        assert converter.warning_lines(text) == [
            "first thing",
            "second thing",
            "bare line",
        ]

    def test_warning_lines_skips_the_reasons_already_in_dropped(self):
        reason = "element 1 (table): data - Field required"
        text = f"[vision_analyzer] dropped {reason}\n[vision_analyzer] 1 of 2 skipped\n"
        assert converter.warning_lines(text, ignore=[reason]) == ["1 of 2 skipped"]


# =========================================================================== #
# app-level wiring
# =========================================================================== #
class TestAppWiring:
    def test_cors_allows_both_dev_origins(self, client):
        assert ALLOWED_ORIGINS == ["http://localhost:5173", "http://localhost:3000"]
        for origin in ALLOWED_ORIGINS:
            response = client.options(
                "/api/convert",
                headers={
                    "Origin": origin,
                    "Access-Control-Request-Method": "POST",
                    "Access-Control-Request-Headers": "content-type,x-anything",
                },
            )
            assert response.status_code == 200
            assert response.headers["access-control-allow-origin"] == origin
            assert "POST" in response.headers["access-control-allow-methods"]

    def test_cors_is_present_on_a_real_response(self, client):
        response = client.get(
            "/api/health", headers={"Origin": "http://localhost:5173"}
        )
        assert response.headers["access-control-allow-origin"] == "http://localhost:5173"

    def test_cors_exposes_the_custom_pptx_headers(self, client):
        response = client.options(
            "/api/download/" + uuid.uuid4().hex,
            headers={
                "Origin": "http://localhost:5173",
                "Access-Control-Request-Method": "GET",
            },
        )
        assert response.status_code == 200

        posted = client.post(
            "/api/build-from-json",
            content=json.dumps(spec_dict()),
            headers={"content-type": "application/json", "Origin": "http://localhost:5173"},
        )
        exposed = posted.headers.get("access-control-expose-headers", "")
        for header in ("X-Element-Counts", "X-Dropped-Count", "Content-Disposition"):
            assert header in exposed

    def test_an_unlisted_origin_is_not_allowed(self, client):
        response = client.get("/api/health", headers={"Origin": "http://evil.example"})
        assert response.status_code == 200
        assert "access-control-allow-origin" not in response.headers

    def test_every_route_is_under_api(self, client):
        schema = client.get("/openapi.json").json()
        paths = {path for path in schema["paths"] if path.startswith("/api")}
        assert paths == set(schema["paths"]) - {"/"}
        assert paths == {
            "/api/health",
            "/api/convert",
            "/api/refine-instructions",
            "/api/preview",
            "/api/build-from-json",
            "/api/download/{file_id}",
        }

    def test_an_unknown_path_uses_the_same_envelope(self, client):
        response = client.get("/api/nope")
        assert response.status_code == 404
        assert envelope(response)["code"] == "not_found"

    def test_a_wrong_method_uses_the_same_envelope(self, client):
        response = client.get("/api/convert")
        assert response.status_code == 405
        assert envelope(response)["code"] == "method_not_allowed"

    def test_the_app_has_a_title_and_version(self):
        assert app.title == "image-to-pptx API"
        assert app.version
        assert app.description

    def test_the_openapi_schema_builds(self, client):
        schema = client.get("/openapi.json")
        assert schema.status_code == 200
        paths = schema.json()["paths"]
        for path in ("/api/health", "/api/convert", "/api/build-from-json"):
            assert path in paths

    def test_the_root_points_at_the_api(self, client):
        body = client.get("/").json()
        assert body["health"] == "/api/health"

    @pytest.mark.parametrize(
        ("status", "detail", "code"),
        [
            (404, "Not Found", "not_found"),
            (405, "Method Not Allowed", "method_not_allowed"),
            (400, None, "bad_request"),
            (400, ["a", "list"], "bad_request"),
            (418, "teapot", "error"),
            (400, {"code": "x", "message": "m", "hint": "h"}, "x"),
            (400, {"code": "x", "message": "m"}, "x"),
        ],
    )
    def test_error_body_always_produces_the_envelope(self, status, detail, code):
        body = error_body(status, detail)
        assert set(body) == {"detail"}
        assert set(body["detail"]) == {"code", "message", "hint"}
        assert body["detail"]["code"] == code
        assert all(isinstance(value, str) for value in body["detail"].values())
        assert body["detail"]["message"]


# =========================================================================== #
# the blocking pipeline must never run on the event loop
# =========================================================================== #
class TestEndpointsDoNotBlockTheEventLoop:
    """``/api/convert`` and ``/api/build-from-json`` are ``async def`` but the
    pipeline they call is blocking: a vision call waits ~10 s on the network and
    ``python-pptx`` is CPU-bound.

    Running that inline would occupy the event loop for the whole conversion, so
    the process would serve exactly one request at a time - a health poll issued
    during a conversion would wait for the model to answer. These tests pin the
    handoff to the threadpool, and they are behavioural: nothing here asserts
    *how* the work is dispatched, only that it does not happen on the loop and
    that two requests really do overlap.
    """

    @staticmethod
    def _loop_state() -> str:
        """``"on the event loop"`` or ``"off the event loop"`` for the CALLER.

        A threadpool worker has no running loop, so ``get_running_loop`` raising
        is the direct signal that this code is not sitting on the loop.
        """
        import asyncio

        try:
            asyncio.get_running_loop()
        except RuntimeError:
            return "off the event loop"
        return "on the event loop"

    def test_convert_runs_the_pipeline_off_the_event_loop(
        self, client, png_bytes, with_key, monkeypatch
    ):
        seen: dict[str, str] = {}

        def _analyze(image_path, **kwargs):
            seen["where"] = self._loop_state()
            return mock_data.sample_spec()

        monkeypatch.setattr(vision_analyzer, "analyze_image", _analyze)

        response = client.post("/api/convert", files=upload(png_bytes))
        assert response.status_code == 200
        assert seen["where"] == "off the event loop"

    def test_the_render_half_of_convert_also_runs_off_the_event_loop(
        self, client, png_bytes, offline_convert, monkeypatch
    ):
        """The vision call is not the only blocking half: rendering the deck is
        CPU-bound, and it must be off the loop too. ``build_deck`` is the step
        the pipeline actually calls, so that is what is instrumented."""
        seen: dict[str, str] = {}
        real_build = converter.build_deck

        def _build(spec, **kwargs):
            seen["where"] = self._loop_state()
            return real_build(spec, **kwargs)

        monkeypatch.setattr(converter, "build_deck", _build)

        response = client.post("/api/convert", files=upload(png_bytes))
        assert response.status_code == 200
        assert seen["where"] == "off the event loop"

    def test_build_from_json_runs_off_the_event_loop(self, client, monkeypatch):
        seen: dict[str, str] = {}
        real_build = converter.build_deck

        def _build(spec, **kwargs):
            seen["where"] = self._loop_state()
            return real_build(spec, **kwargs)

        monkeypatch.setattr(converter, "build_deck", _build)

        response = client.post("/api/build-from-json", json=spec_dict())
        assert response.status_code == 200
        assert seen["where"] == "off the event loop"

    def test_preview_runs_off_the_event_loop(self, client, monkeypatch):
        seen: dict[str, str] = {}
        real_build = converter.build_deck

        def _build(spec, **kwargs):
            seen["where"] = self._loop_state()
            return real_build(spec, **kwargs)

        monkeypatch.setattr(converter, "build_deck", _build)

        response = client.post("/api/preview", json=spec_dict())
        assert response.status_code == 200
        assert seen["where"] == "off the event loop"

    def test_two_conversions_are_genuinely_concurrent(
        self, client, png_bytes, with_key, monkeypatch
    ):
        """A deterministic concurrency proof, with no timing threshold.

        Both requests must be inside ``analyze_image`` at the same moment to get
        past the barrier. If the pipeline ran on the event loop the second
        request could not start until the first returned, so the barrier would
        time out - which is exactly the stall this guards against.
        """
        import threading

        barrier = threading.Barrier(2, timeout=10)
        # Idents, not names: anyio gives every worker the same thread NAME.
        reached: list[int] = []

        def _analyze(image_path, **kwargs):
            reached.append(threading.get_ident())
            barrier.wait()  # BrokenBarrierError on timeout -> a 500 below
            return mock_data.sample_spec()

        monkeypatch.setattr(vision_analyzer, "analyze_image", _analyze)

        statuses: list[int] = []

        def _fire():
            statuses.append(
                client.post("/api/convert", files=upload(png_bytes)).status_code
            )

        threads = [threading.Thread(target=_fire) for _ in range(2)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(timeout=30)

        assert statuses == [200, 200], (
            "both conversions should succeed; a non-200 means the barrier broke, "
            "i.e. the two requests never overlapped"
        )
        assert len(reached) == 2
        assert len(set(reached)) == 2, (
            f"both conversions ran on the same thread ({reached}) - the pipeline "
            "is not being handed to the threadpool"
        )

    def test_health_is_not_blocked_by_an_in_flight_conversion(
        self, client, png_bytes, with_key, monkeypatch
    ):
        """The contract calls /api/health a report that "never fails". It is
        only usable as a liveness probe if it also never QUEUES behind a
        conversion, which is what an inline vision call would make it do."""
        import threading

        release = threading.Event()
        health_done = threading.Event()

        def _analyze(image_path, **kwargs):
            # Hold the conversion open until /api/health has answered.
            assert health_done.wait(timeout=10), (
                "/api/health did not answer while a conversion was in flight - "
                "the event loop is blocked by the pipeline"
            )
            release.set()
            return mock_data.sample_spec()

        monkeypatch.setattr(vision_analyzer, "analyze_image", _analyze)

        convert_status: list[int] = []

        def _convert():
            convert_status.append(
                client.post("/api/convert", files=upload(png_bytes)).status_code
            )

        worker = threading.Thread(target=_convert)
        worker.start()
        try:
            health = client.get("/api/health")
            assert health.status_code == 200
            assert health.json()["status"] == "ok"
        finally:
            health_done.set()
            worker.join(timeout=30)

        assert release.is_set()
        assert convert_status == [200]


# =========================================================================== #
# the response body never carries credential material
# =========================================================================== #
class TestNoSecretsInResponses:
    """Regression tests for a real leak: when LiteLLM rejects the configured key
    its 401 body says ``Received API Key = sk-...<last 4>, Key Hash (Token)
    =<sha256 of the whole key>``. The OpenAI SDK puts that body in its exception
    message, ``vision_analyzer`` wraps it, and the 502 ``analysis_failed`` hint
    used to relay ``str(exc)`` verbatim to any unauthenticated caller - the
    suffix narrows a brute force and the digest confirms a candidate key
    offline. Nothing on any path may publish either.
    """

    #: A stand-in for a rejected key. Obviously fake (see the module docstring),
    #: but shaped exactly like the thing that must not escape.
    LEAKY_KEY = "sk-NOT-A-REAL-KEY-0000000000000000"

    @property
    def key_digest(self) -> str:
        import hashlib

        return hashlib.sha256(self.LEAKY_KEY.encode()).hexdigest()

    def upstream_401_text(self) -> str:
        """What the proxy's 401 actually looks like, reproduced verbatim."""
        return (
            f"Vision request to model 'gemini-flash' failed.\n"
            f"  first attempt (with response_format): AuthenticationError: Error "
            f"code: 401 - {{'error': {{'message': 'Authentication Error, Invalid "
            f"proxy server token passed. Received API Key = "
            f"{self.LEAKY_KEY[:3]}...{self.LEAKY_KEY[-4:]}, Key Hash (Token) "
            f"={self.key_digest}. Unable to find token in cache', 'type': "
            f"'token_not_found_in_db', 'param': 'key', 'code': '401'}}}}\n"
            f"  retry (plain): AuthenticationError: Error code: 401 - Received "
            f"API Key = {self.LEAKY_KEY[:3]}...{self.LEAKY_KEY[-4:]}, Key Hash "
            f"(Token) ={self.key_digest}.\n"
            f"Check that the LiteLLM proxy at http://localhost:4000/ is running."
        )

    def assert_clean(self, response, *, text: str | None = None) -> None:
        """No key material anywhere in the response.

        ``text`` narrows the "no hex digest at all" sweep to the part of the
        body that relays upstream words, because a 200 legitimately carries a
        ``download_id`` (a uuid4 hex this service generated itself).
        """
        body = response.text
        assert "sk-" not in body, "a key (or its suffix) reached the client"
        assert self.key_digest not in body, "sha256(key) reached the client"
        assert self.LEAKY_KEY not in body
        import re

        assert not re.search(
            r"\b[0-9a-fA-F]{32,}\b", body if text is None else text
        ), "a digest reached the client"

    def test_a_rejected_key_is_not_relayed_by_analysis_failed(
        self, client, png_bytes, with_key, monkeypatch
    ):
        def _rejected(*args, **kwargs):
            raise vision_analyzer.VisionAnalysisError(self.upstream_401_text())

        monkeypatch.setattr(vision_analyzer, "analyze_image", _rejected)
        response = client.post("/api/convert", files=upload(png_bytes))

        assert response.status_code == 502
        detail = envelope(response)
        assert detail["code"] == "analysis_failed"
        self.assert_clean(response)
        # Still actionable: the model name and the proxy advice survive.
        assert "gemini-flash" in detail["hint"]
        assert "server log" in detail["hint"].lower()

    def test_the_unredacted_text_is_logged_instead_of_published(
        self, client, png_bytes, with_key, monkeypatch, caplog
    ):
        """Scrubbing must not lose the diagnosis, only relocate it."""

        def _rejected(*args, **kwargs):
            raise vision_analyzer.VisionAnalysisError(self.upstream_401_text())

        monkeypatch.setattr(vision_analyzer, "analyze_image", _rejected)
        with caplog.at_level("WARNING", logger="app.services.converter"):
            response = client.post("/api/convert", files=upload(png_bytes))

        assert response.status_code == 502
        self.assert_clean(response)
        assert self.key_digest in caplog.text, "the operator lost the upstream error"

    def test_the_log_line_never_lands_in_the_warnings_array(self, caplog):
        """The operator's copy must leave the process as a LOG, not as a field.

        ``failure_from_exception`` runs inside ``capture_stderr`` (the pipeline
        wraps the whole conversion in it) and python's default logging handler
        resolves ``sys.stderr`` when it emits, so a naive ``logger.warning``
        wrote the unredacted error into the request's warning buffer - a
        RESPONSE field - and it was then discarded on the failure path, leaving
        nobody at all with the upstream error.
        """
        with caplog.at_level("WARNING", logger="app.services.converter"):
            with converter.capture_stderr() as buffer:
                converter.failure_from_exception(
                    vision_analyzer.VisionAnalysisError(self.upstream_401_text()),
                    context="convert",
                )
                captured_by_the_request = buffer.getvalue()

        assert captured_by_the_request == "", (
            "the upstream error was captured as a response warning"
        )
        assert self.key_digest in caplog.text, "the operator lost it entirely"

    def test_an_operator_diagnostic_bypasses_the_request_buffer(self):
        """The mechanism the test above depends on, asserted on its own."""
        import sys

        with converter.capture_stderr() as buffer:
            with converter._uncaptured_stderr():
                print("operator-only diagnostic", file=sys.stderr)
            print("[pptx_builder] warning: meant for the client", file=sys.stderr)

        assert "operator-only" not in buffer.getvalue()
        assert "meant for the client" in buffer.getvalue()

    def test_the_configured_key_is_removed_even_in_an_unknown_shape(
        self, client, png_bytes, with_key, monkeypatch
    ):
        """The patterns are a net, not the guarantee: the key's own value is
        removed literally, so a proxy error worded some new way cannot carry
        it out either."""

        def _rejected(*args, **kwargs):
            raise vision_analyzer.VisionAnalysisError(
                f"the upstream said <<{with_key}>> in prose nobody predicted"
            )

        monkeypatch.setattr(vision_analyzer, "analyze_image", _rejected)
        response = client.post("/api/convert", files=upload(png_bytes))
        assert response.status_code == 502
        assert with_key not in response.text

    def test_a_retry_warning_cannot_carry_key_material_into_a_200(
        self, client, png_bytes, with_key, monkeypatch
    ):
        """The second leak channel: when the response_format attempt fails and
        the retry succeeds, vision_analyzer prints the first failure - proxy
        error body and all - on stderr, and the contract returns that array in
        a 200 ``warnings``."""
        import sys

        def _warn_then_succeed(*args, **kwargs):
            print(
                f"[vision_analyzer] model 'gemini-flash' failed "
                f"(AuthenticationError: Received API Key = sk-...0000, Key Hash "
                f"(Token) ={self.key_digest}); retrying without response_format",
                file=sys.stderr,
            )
            return mock_data.sample_spec()

        monkeypatch.setattr(vision_analyzer, "analyze_image", _warn_then_succeed)
        response = client.post("/api/convert", files=upload(png_bytes))

        assert response.status_code == 200
        warnings = response.json()["warnings"]
        self.assert_clean(response, text=json.dumps(warnings))
        assert any("retrying without response_format" in line for line in warnings), (
            "the warning itself must survive - only the credential is dropped"
        )

    def test_health_never_carries_key_material(self, client, with_key):
        response = client.get("/api/health")
        assert response.status_code == 200
        assert response.json()["proxy"]["api_key_configured"] is True
        assert with_key not in response.text

    @pytest.mark.parametrize(
        "text",
        [
            "Received API Key = sk-abcd1234efgh, so no",
            "Key Hash (Token) =0123456789abcdef0123456789abcdef0123456789abcdef0123456789abcdef",
            "authorization: Bearer sk-livekey-abcdef123456",
            "token=abcdef0123456789abcdef0123456789",
            "the key is sk-0000111122223333",
        ],
    )
    def test_scrub_secrets_drops_every_known_shape(self, text):
        scrubbed = converter.scrub_secrets(text)
        assert "sk-" not in scrubbed
        assert "[redacted]" in scrubbed
        import re

        assert not re.search(r"\b[0-9a-fA-F]{32,}\b", scrubbed)

    def test_scrub_secrets_keeps_ordinary_diagnostics(self):
        text = (
            "Vision request to model 'gemini-flash' failed.\n"
            "  bbox (505, 601, 461, 5): height is below the 0.05in minimum\n"
            "  element 1 (table): data - Field required"
        )
        assert converter.scrub_secrets(text) == text

    def test_the_missing_key_hint_still_names_the_variable_and_the_export(self):
        """The scrub may not eat the one hint that has to mention the key's NAME."""
        failure = converter.failure_from_exception(
            config.MissingApiKeyError("no key"), context="convert"
        )
        assert "LITELLM_API_KEY" in failure.hint
        assert "export LITELLM_API_KEY=" in failure.hint
        assert "/api/build-from-json" in failure.hint

    def test_a_server_side_debug_flag_can_restore_the_raw_text(self, monkeypatch):
        """Escape hatch for an operator debugging the proxy: off by default, and
        only settable in the server's own environment."""
        raw = f"Received API Key = {self.LEAKY_KEY}"
        assert self.LEAKY_KEY not in converter.scrub_secrets(raw)
        monkeypatch.setenv(converter.DEBUG_UPSTREAM_ENV_VAR, "1")
        assert converter.scrub_secrets(raw) == raw


# =========================================================================== #
# POST /api/build-from-json: the body is bounded, and every malformed body is
# still the endpoint's own 400
# =========================================================================== #
class TestJsonBodyLimits:
    def post_raw(self, client, body, **headers):
        return client.post(
            "/api/build-from-json",
            content=body,
            headers={"content-type": "application/json", **headers},
        )

    def test_the_cap_is_declared_and_generous_but_finite(self):
        # A SlideSpec is kilobytes; the cap only has to stop a body that would
        # cost the process multiples of its own size to decode.
        assert 256 * 1024 <= converter.MAX_JSON_BODY_BYTES <= 8 * 1024 * 1024
        assert converter.MAX_JSON_BODY_BYTES < converter.MAX_UPLOAD_BYTES

    def test_a_body_past_the_cap_is_refused(self, client):
        filler = "Q" * (converter.MAX_JSON_BODY_BYTES + 1024)
        body = '{"background_color":"' + filler + '","elements":[]}'
        response = self.post_raw(client, body)
        assert response.status_code == 400
        detail = envelope(response)
        assert detail["code"] == "invalid_spec"
        assert "limit" in detail["message"]

    def test_a_body_at_the_cap_is_still_accepted(self, client):
        """The cap may not clip a legitimate spec: a body just under it works."""
        spec = spec_dict()
        pad = converter.MAX_JSON_BODY_BYTES - len(json.dumps(spec)) - 64
        spec["elements"].append(
            {
                "type": "text",
                "bbox": [0, 0, 100, 100],
                "content": "x" * max(pad, 1),
            }
        )
        body = json.dumps(spec)
        assert len(body) < converter.MAX_JSON_BODY_BYTES
        response = self.post_raw(client, body)
        assert response.status_code == 200
        assert response.content[:2] == b"PK"

    def test_a_chunked_body_is_refused_too(self, client):
        """No Content-Length to check, so the count has to happen on the stream."""
        chunk = b"Q" * (256 * 1024)
        total = (converter.MAX_JSON_BODY_BYTES // len(chunk)) + 40

        def body():
            yield b'{"background_color":"'
            for _ in range(total):
                yield chunk
            yield b'","elements":[]}'

        response = client.post(
            "/api/build-from-json",
            content=body(),
            headers={"content-type": "application/json"},
        )
        assert response.status_code == 400
        assert envelope(response)["code"] == "invalid_spec"

    def test_the_reader_stops_pulling_chunks_at_the_cap(self):
        """Refusing must mean reading STOPS, not just that the answer is a 400.

        Asserted on the reader itself because TestClient's ASGI transport
        buffers the request body before the app ever runs, which would hide the
        property this cap exists for.
        """
        import asyncio

        from fastapi import HTTPException

        from app.api import routes

        chunk = b"Q" * (256 * 1024)
        pulled = {"chunks": 0}
        total = (converter.MAX_JSON_BODY_BYTES // len(chunk)) + 40

        class _ChunkedRequest:
            headers: dict[str, str] = {}  # chunked: no content-length at all

            async def stream(self):
                for _ in range(total):
                    pulled["chunks"] += 1
                    yield chunk

        with pytest.raises(HTTPException) as caught:
            asyncio.run(routes._read_json_body_capped(_ChunkedRequest()))
        assert caught.value.status_code == 400
        assert caught.value.detail["code"] == "invalid_spec"
        assert pulled["chunks"] <= (converter.MAX_JSON_BODY_BYTES // len(chunk)) + 1
        assert pulled["chunks"] < total

    def test_a_declared_oversize_body_is_refused_before_it_is_read(self):
        """An honest Content-Length is refused without reading a single byte."""
        import asyncio

        from fastapi import HTTPException

        from app.api import routes

        class _DeclaredRequest:
            headers = {"content-length": str(converter.MAX_JSON_BODY_BYTES * 8)}

            async def stream(self):  # pragma: no cover - a read is the failure
                raise AssertionError("the body was read despite Content-Length")
                yield b""

        with pytest.raises(HTTPException) as caught:
            asyncio.run(routes._read_json_body_capped(_DeclaredRequest()))
        assert caught.value.status_code == 400
        assert caught.value.detail["code"] == "invalid_spec"

    @pytest.mark.parametrize("depth", [10_000, 200_000])
    def test_a_deeply_nested_body_is_invalid_spec_not_a_500(self, client, depth):
        """CPython answers a deeply nested document with RecursionError, which is
        NOT a ValueError: guarding only ValueError let it escape the route and
        leave as a 500 ``internal_error``, a code this endpoint's contract does
        not list. Caller garbage is ``invalid_spec``, however it is malformed."""
        body = b"[" * depth + b"]" * depth
        assert len(body) < converter.MAX_JSON_BODY_BYTES
        response = self.post_raw(client, body)
        assert response.status_code == 400
        detail = envelope(response)
        assert detail["code"] == "invalid_spec"
        assert "nested" in detail["message"].lower()

    def test_the_server_survives_a_deeply_nested_body(self, client):
        self.post_raw(client, b"[" * 100_000 + b"]" * 100_000)
        assert client.get("/api/health").status_code == 200
        assert self.post_raw(client, json.dumps(spec_dict())).status_code == 200

    def test_a_nesting_depth_a_spec_could_plausibly_use_still_validates(self, client):
        """The fix may not turn "deep" into "rejected": a body that parses must
        keep reaching the schema."""
        response = self.post_raw(client, b"[" * 400 + b"]" * 400)
        assert response.status_code == 400
        assert envelope(response)["code"] == "invalid_spec"

    def test_parse_json_body_maps_every_failure_onto_invalid_spec(self):
        cases = [
            b"",
            b"   ",
            b"{not json at all",
            b"[" * 20_000 + b"]" * 20_000,
            b'{"a":' + b"Q" * (converter.MAX_JSON_BODY_BYTES + 1),
        ]
        for raw in cases:
            with pytest.raises(ConversionFailure) as caught:
                converter.parse_json_body(raw)
            assert (caught.value.status_code, caught.value.code) == (400, "invalid_spec")

    def test_parse_json_body_still_parses_what_it_should(self):
        assert converter.parse_json_body(b'{"elements": []}') == {"elements": []}
        assert converter.parse_json_body('[{"type": "text"}]') == [{"type": "text"}]


# =========================================================================== #
# a 500 is readable from the browser origins the contract allows
# =========================================================================== #
class TestErrorEnvelopeIsReadableCrossOrigin:
    """``ServerErrorMiddleware`` (where starlette installs the ``Exception``
    handler) sits OUTSIDE ``CORSMiddleware``, so a 500 produced there carries no
    ``Access-Control-Allow-Origin`` and the browser rejects it as a network
    failure - exactly when the user most needs the hint. The envelope is
    produced one layer further in now, and these tests pin the header onto every
    status the app can answer with, 500 included.
    """

    ORIGIN = "http://localhost:5173"

    def test_a_500_from_convert_carries_the_cors_header(
        self, client, png_bytes, monkeypatch
    ):
        def _boom(*args, **kwargs):
            raise RuntimeError("secret internal detail at /srv/app/thing.py:42")

        monkeypatch.setattr(converter, "convert_upload", _boom)
        response = client.post(
            "/api/convert",
            files=upload(png_bytes),
            headers={"Origin": self.ORIGIN},
        )
        assert response.status_code == 500
        assert response.headers.get("access-control-allow-origin") == self.ORIGIN
        detail = envelope(response)
        assert detail["code"] == "internal_error"
        assert "RuntimeError" in detail["message"]
        # The envelope is still saying nothing an attacker can use.
        assert "secret internal detail" not in response.text
        assert "Traceback" not in response.text
        assert ".py" not in response.text

    def test_a_500_from_build_from_json_carries_the_cors_header(
        self, client, monkeypatch
    ):
        def _boom(payload):
            raise RuntimeError("nope")

        monkeypatch.setattr(converter, "build_from_payload", _boom)
        response = client.post(
            "/api/build-from-json",
            content=json.dumps(spec_dict()),
            headers={"content-type": "application/json", "Origin": self.ORIGIN},
        )
        assert response.status_code == 500
        assert response.headers.get("access-control-allow-origin") == self.ORIGIN
        assert envelope(response)["code"] == "internal_error"

    @pytest.mark.parametrize("origin", ALLOWED_ORIGINS)
    def test_every_allowed_origin_can_read_a_500(
        self, client, png_bytes, monkeypatch, origin
    ):
        monkeypatch.setattr(
            converter,
            "convert_upload",
            lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom")),
        )
        response = client.post(
            "/api/convert",
            files=upload(png_bytes),
            headers={"Origin": origin},
        )
        assert response.status_code == 500
        assert response.headers.get("access-control-allow-origin") == origin

    def test_a_500_does_not_hand_cors_to_an_unlisted_origin(
        self, client, png_bytes, monkeypatch
    ):
        """Moving the envelope inside the CORS layer must not make the layer
        more permissive than it was."""
        monkeypatch.setattr(
            converter,
            "convert_upload",
            lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom")),
        )
        response = client.post(
            "/api/convert",
            files=upload(png_bytes),
            headers={"Origin": "http://evil.example"},
        )
        assert response.status_code == 500
        assert "access-control-allow-origin" not in response.headers

    @pytest.mark.parametrize(
        ("path", "kwargs", "expected"),
        [
            ("/api/health", {}, 200),
            ("/api/nope", {}, 404),
            ("/api/download/not-a-uuid", {}, 404),
        ],
    )
    def test_the_other_statuses_keep_the_header(self, client, path, kwargs, expected):
        response = client.get(path, headers={"Origin": self.ORIGIN}, **kwargs)
        assert response.status_code == expected
        assert response.headers.get("access-control-allow-origin") == self.ORIGIN

    def test_the_outer_handler_is_still_installed_as_the_net_behind_it(self):
        """The middleware handles what the app raises; the handler outside CORS
        stays for what the CORS layer itself could raise."""
        from fastapi.middleware.cors import CORSMiddleware

        from app import main

        assert Exception in app.exception_handlers
        assert any(
            middleware.cls is main.ErrorEnvelopeMiddleware
            for middleware in app.user_middleware
        )
        # ...and it is INSIDE the CORS layer, i.e. later in the list.
        classes = [middleware.cls for middleware in app.user_middleware]
        assert classes.index(main.ErrorEnvelopeMiddleware) > classes.index(
            CORSMiddleware
        )


# =========================================================================== #
# the unsupported_extension message has to be actionable
# =========================================================================== #
class TestUnsupportedExtensionMessage:
    @pytest.mark.parametrize("name", [".png", ".jpg", ".jpeg", ".webp", ".PNG"])
    def test_a_dotfile_name_is_never_described_as_its_own_extension(
        self, client, png_bytes, name
    ):
        """``Path(".png").suffix`` is '' - POSIX reads a leading-dot name as a
        name, not a suffix - and quoting the filename there produced
        "Unsupported image type '.png'. Supported extensions: ..., .png, ...":
        a sentence that rejects and allows the same thing, leaving the user
        nothing to act on."""
        response = client.post(
            "/api/convert",
            files=upload(png_bytes, name=name),
        )
        assert response.status_code == 400
        detail = envelope(response)
        assert detail["code"] == "unsupported_extension"

        rejected = detail["message"].split("Supported extensions:")[0]
        assert name.lower() not in rejected.lower(), (
            "the message rejects the very extension it goes on to allow"
        )
        assert "<no extension>" in rejected
        # The contract still requires the whole supported list in the message.
        for extension in sorted(config.SUPPORTED_EXTENSIONS):
            assert extension in detail["message"]
        # ...and the hint says what to actually do about it.
        assert "slide.png" in detail["hint"]

    def test_a_real_unsupported_extension_is_still_named(self, client, png_bytes):
        response = client.post(
            "/api/convert",
            files=upload(png_bytes, name="scan.tiff"),
        )
        assert response.status_code == 400
        assert "'.tiff'" in envelope(response)["message"]

    def test_no_filename_at_all_is_still_reported_as_such(self):
        with pytest.raises(ConversionFailure) as caught:
            converter.validate_upload_bytes(b"\x89PNG", None)
        assert caught.value.code == "unsupported_extension"
        assert "<no filename>" in caught.value.message

    @pytest.mark.parametrize("filename", [".png", ".jpg", "noext", "..png", "a.png."])
    def test_a_name_that_cannot_be_accepted_is_described_accurately(self, filename):
        """Whatever the reason, the message names a type the user can compare
        with the supported list - never a filename dressed as an extension."""
        extension = converter.upload_extension(filename)
        if extension in converter.supported_extensions():
            return
        with pytest.raises(ConversionFailure) as caught:
            converter.validate_upload_bytes(b"\x89PNG", filename)
        message = caught.value.message.split("Supported extensions:")[0]
        assert filename not in message
