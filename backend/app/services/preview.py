"""Describe a BUILT deck as a render-ready payload the UI can draw.

The whole point of this module is where it reads from: it walks the shapes of a
real :class:`pptx.presentation.Presentation` and reports the geometry and
styling python-pptx **actually** produced. It never looks at the ``SlideSpec``.

That is deliberate, and it is the only design that can be correct. Re-deriving
the preview from the spec would mean a second implementation of
``pptx_builder.scale_bbox`` - and the moment the two disagreed, the user would
be shown a picture of a slide that does not exist. Reading the built deck means
the preview shows, for free and by construction:

* the min-dimension floor (``MIN_DIMENSION_IN``), so a 5-unit-tall band is
  previewed at the height it was actually enlarged to;
* the off-slide clipping, so a bbox that overshot 1000 is previewed where it
  really landed;
* the cleaned text - control characters dropped, CRLF split into paragraphs,
  lone surrogates replaced - because the runs are read back out of the deck;
* the z-order, because ``slide.shapes`` is in document order, which *is* the
  z-order, back to front.

Output units
------------
Geometry is emitted as **percentages of the slide** so the UI can scale to any
width with no arithmetic, and font sizes in **points** so text can be sized
against ``slide.height_pt`` (540pt for the 7.5in slide). Every numeric field is
a JSON number. Percentages are deliberately NOT clamped: a shape may legitimately
sit partly off the slide, and clipping it is the UI's job, not this module's.

Colours
-------
Only an explicit RGB colour is reported. A theme colour, an inherited colour or
no fill at all comes back as ``None`` - never a guess and never an exception.
This matters more than it looks: on python-pptx 1.0.2 (verified)
``fill.fore_color.rgb`` RAISES rather than returning a falsy value - ``TypeError``
when the fill is not solid, ``AttributeError`` when the colour is a scheme
colour - so every colour read here is guarded.

Robustness
----------
Every accessor is wrapped. A shape this module cannot make sense of degrades to
``{"kind": "unsupported"}`` with its geometry rather than raising, and a picture
shape - which this product never emits, since nothing is ever flattened - would
do the same. The caller (``converter.preview_for``) additionally treats any
escape from here as "no preview, plus a warning", so a preview can never fail a
conversion.
"""

from __future__ import annotations

from typing import Any

from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from .. import config

__all__ = ["build_preview"]

#: python-pptx auto-shape preset -> the name the preview contract uses.
#: Anything else is reported as ``"other"``; the UI draws it as a plain box.
_AUTOSHAPE_NAMES: dict[str, str] = {
    "RECTANGLE": "rectangle",
    "ROUNDED_RECTANGLE": "rounded_rectangle",
    "OVAL": "oval",
}

#: Paragraph alignment -> the contract's three values. ``JUSTIFY`` and friends
#: have no preview representation, so they read as the default.
_ALIGN_NAMES: dict[str, str] = {
    "LEFT": "left",
    "CENTER": "center",
    "RIGHT": "right",
}

#: The contract's default alignment, used whenever none is set in the file.
_DEFAULT_ALIGN = "left"

#: Percentages are rounded here: enough precision that a 1px feature on a 4K
#: screen still lands right, few enough digits that the JSON stays readable.
_PCT_DIGITS = 4

#: Point sizes are rounded to the precision PowerPoint itself stores (1/100 pt).
_PT_DIGITS = 2

#: Shape types that would mean part of the image was flattened into the deck.
#: The builder emits none, ever - this exists so that if one somehow appeared,
#: the preview says "unsupported" instead of drawing a box where a picture is.
_PICTURE_TYPES = (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE)


# --------------------------------------------------------------------------- #
# guarded readers - none of these may raise, whatever the file contains
# --------------------------------------------------------------------------- #
def _pct(value: Any, total: Any) -> float:
    """``value`` as a percentage of ``total``, as a JSON number."""
    try:
        divisor = float(total)
        if divisor <= 0:
            return 0.0
        return round(float(value) / divisor * 100.0, _PCT_DIGITS)
    except (TypeError, ValueError, ZeroDivisionError):
        return 0.0


def _rgb(color: Any) -> str | None:
    """``"#RRGGBB"`` when the colour is an explicit RGB, else ``None``.

    Reading ``.rgb`` on anything other than an RGB colour raises on python-pptx
    1.0.2, so the TYPE is checked first and the read is guarded anyway.
    """
    try:
        if color is None or color.type != MSO_COLOR_TYPE.RGB:
            return None
        return f"#{str(color.rgb).upper()}"
    except Exception:  # noqa: BLE001 - a colour may never break a preview
        return None


def _solid_fill(fill: Any) -> str | None:
    """The fill's explicit RGB colour, or ``None`` for no-fill / theme / other.

    ``None`` is what the UI draws as see-through, which is exactly right for
    both "the spec asked for no fill" and "this colour is not an RGB we can
    show" - in neither case may the preview invent a colour.
    """
    try:
        if fill is None or fill.type != MSO_FILL.SOLID:
            return None
        return _rgb(fill.fore_color)
    except Exception:  # noqa: BLE001
        return None


def _line_color(line: Any) -> str | None:
    """The outline's explicit RGB colour, or ``None`` when there is no outline."""
    try:
        if line is None or line.fill.type != MSO_FILL.SOLID:
            return None
        return _rgb(line.color)
    except Exception:  # noqa: BLE001
        return None


def _line_width_pt(line: Any) -> float | None:
    """The outline width in points, or ``None`` when there is none.

    An unset width reads back as ``Emu(0)`` rather than ``None``, and a 0pt
    outline is not an outline, so zero is reported as ``None`` too.
    """
    try:
        width = None if line is None else line.width
        if width is None:
            return None
        points = round(float(Emu(int(width)).pt), _PT_DIGITS)
        return points if points > 0 else None
    except Exception:  # noqa: BLE001
        return None


def _font_size_pt(fonts: list[Any], *, default: float) -> float:
    """The first explicit size among ``fonts``, else the builder's default.

    The fallback is what keeps the UI from having to guess: a run that carries
    no size inherits one from PowerPoint, and the value it would inherit is the
    same one ``pptx_builder`` passes to ``_positive(...)``, so the preview
    reports that rather than ``null``.
    """
    for font in fonts:
        try:
            size = font.size
        except Exception:  # noqa: BLE001
            continue
        if size is None:
            continue
        try:
            return round(float(size.pt), _PT_DIGITS)
        except (AttributeError, TypeError, ValueError):
            continue
    return float(default)


def _font_color(fonts: list[Any]) -> str | None:
    """The first explicit RGB colour among ``fonts``, else ``None``."""
    for font in fonts:
        try:
            color = font.color
        except Exception:  # noqa: BLE001
            continue
        found = _rgb(color)
        if found is not None:
            return found
    return None


def _bold(fonts: list[Any]) -> bool:
    """The first explicit bold state among ``fonts``, else ``False``.

    ``None`` means "inherit", which the builder never writes for text it owns -
    it writes ``b="0"`` on purpose - so a ``None`` here really is an absence.
    """
    for font in fonts:
        try:
            value = font.bold
        except Exception:  # noqa: BLE001
            continue
        if value is not None:
            return bool(value)
    return False


def _align(paragraph: Any, default: str = _DEFAULT_ALIGN) -> str:
    """The paragraph's alignment as ``"left"``/``"center"``/``"right"``."""
    try:
        alignment = None if paragraph is None else paragraph.alignment
        if alignment is None:
            return default
        name = getattr(alignment, "name", None) or str(alignment)
        return _ALIGN_NAMES.get(str(name).upper(), default)
    except Exception:  # noqa: BLE001
        return default


def _autoshape(shape: Any) -> str:
    """The auto-shape preset as the contract names it, else ``"other"``."""
    try:
        preset = shape.auto_shape_type
        if preset is None:
            return "other"
        name = getattr(preset, "name", None) or str(preset)
        return _AUTOSHAPE_NAMES.get(str(name).upper(), "other")
    except Exception:  # noqa: BLE001
        return "other"


def _shape_type(shape: Any) -> Any:
    try:
        return shape.shape_type
    except Exception:  # noqa: BLE001
        return None


def _fonts_of(paragraph: Any) -> list[Any]:
    """The fonts to consult for one paragraph, most specific first.

    A run's own font wins over the paragraph's, which is the order PowerPoint
    resolves them in. An empty paragraph has no runs at all - the builder styles
    the paragraph font precisely so that case still carries the right size - so
    the paragraph font is always kept as the fallback.
    """
    fonts: list[Any] = []
    try:
        runs = list(paragraph.runs)
    except Exception:  # noqa: BLE001
        runs = []
    if runs:
        try:
            fonts.append(runs[0].font)
        except Exception:  # noqa: BLE001
            pass
    try:
        fonts.append(paragraph.font)
    except Exception:  # noqa: BLE001
        pass
    return fonts


def _paragraph_text(paragraph: Any) -> str:
    """The paragraph's text, read back out of its runs."""
    try:
        runs = list(paragraph.runs)
    except Exception:  # noqa: BLE001
        runs = []
    if runs:
        return "".join(str(run.text or "") for run in runs)
    try:
        return str(paragraph.text or "")
    except Exception:  # noqa: BLE001
        return ""


def _track_pcts(sizes: list[int], fallback_total: Any) -> list[float]:
    """Column widths / row heights as percentages of the TABLE's own extent.

    The denominator is the sum of the tracks rather than the frame's width or
    height, so the list always sums to ~100 and the UI's grid fills its box
    exactly. ``pptx_builder`` distributes the frame evenly with integer EMU
    division, so the two denominators differ by at most a few EMU anyway - far
    below one screen pixel - and the frame extent is kept only as the fallback
    for a degenerate table whose tracks are all zero.
    """
    total = sum(sizes)
    if total <= 0:
        try:
            total = int(fallback_total or 0)
        except (TypeError, ValueError):
            total = 0
    if total <= 0:
        return [0.0 for _ in sizes]
    return [_pct(size, total) for size in sizes]


# --------------------------------------------------------------------------- #
# per-shape payloads
# --------------------------------------------------------------------------- #
def _geometry(shape: Any, slide_width: int, slide_height: int) -> dict[str, float]:
    """The shape's box as percentages of the slide.

    x and width scale by the slide's WIDTH, y and height by its HEIGHT. On a
    16:9 slide those differ by 1.78x, so swapping them is not a rounding error -
    it is a visibly wrong picture. ``tests/test_preview.py`` asserts the pairing
    against a deliberately non-square box for exactly that reason.
    """
    return {
        "left_pct": _pct(shape.left, slide_width),
        "top_pct": _pct(shape.top, slide_height),
        "width_pct": _pct(shape.width, slide_width),
        "height_pct": _pct(shape.height, slide_height),
    }


def _text_payload(shape: Any, geometry: dict[str, float]) -> dict[str, Any]:
    """A text box: one entry per paragraph, each with its own styling."""
    paragraphs: list[dict[str, Any]] = []
    try:
        frame_paragraphs = list(shape.text_frame.paragraphs)
    except Exception:  # noqa: BLE001
        frame_paragraphs = []

    for paragraph in frame_paragraphs:
        fonts = _fonts_of(paragraph)
        paragraphs.append(
            {
                "text": _paragraph_text(paragraph),
                "font_size_pt": _font_size_pt(
                    fonts, default=config.DEFAULT_TEXT_FONT_SIZE
                ),
                "font_color": _font_color(fonts),
                "bold": _bold(fonts),
                "align": _align(paragraph),
            }
        )

    return {
        "kind": "text",
        **geometry,
        # The block-level alignment the UI can apply before it reads the
        # paragraphs; the builder writes the same value on every one of them.
        "align": paragraphs[0]["align"] if paragraphs else _DEFAULT_ALIGN,
        "paragraphs": paragraphs,
    }


def _cell_payload(cell: Any) -> dict[str, Any]:
    """One table cell: its text, its fill and the font its runs really carry."""
    try:
        paragraphs = list(cell.text_frame.paragraphs)
    except Exception:  # noqa: BLE001
        paragraphs = []
    first = paragraphs[0] if paragraphs else None
    fonts = _fonts_of(first) if first is not None else []

    try:
        text = str(cell.text or "")
    except Exception:  # noqa: BLE001
        text = ""

    return {
        "text": text,
        "fill_color": _solid_fill(getattr(cell, "fill", None)),
        "font_size_pt": _font_size_pt(fonts, default=config.DEFAULT_TABLE_FONT_SIZE),
        "font_color": _font_color(fonts),
        "bold": _bold(fonts),
        "align": _align(first),
    }


def _table_payload(shape: Any, geometry: dict[str, float]) -> dict[str, Any]:
    """A native table: its track sizes and every cell, in row-major order."""
    table = shape.table

    col_widths = [int(column.width or 0) for column in table.columns]
    row_heights = [int(row.height or 0) for row in table.rows]

    rows: list[list[dict[str, Any]]] = []
    for row_index in range(len(row_heights)):
        cells: list[dict[str, Any]] = []
        for column_index in range(len(col_widths)):
            cells.append(_cell_payload(table.cell(row_index, column_index)))
        rows.append(cells)

    return {
        "kind": "table",
        **geometry,
        "col_widths_pct": _track_pcts(col_widths, getattr(shape, "width", 0)),
        "row_heights_pct": _track_pcts(row_heights, getattr(shape, "height", 0)),
        "rows": rows,
    }


def _shape_payload(shape: Any, geometry: dict[str, float]) -> dict[str, Any]:
    """An auto shape: its preset, its fill and its outline."""
    line = getattr(shape, "line", None)
    return {
        "kind": "shape",
        "autoshape": _autoshape(shape),
        **geometry,
        "fill_color": _solid_fill(getattr(shape, "fill", None)),
        "line_color": _line_color(line),
        "line_width_pt": _line_width_pt(line),
    }


def _one_shape(shape: Any, slide_width: int, slide_height: int) -> dict[str, Any]:
    """Dispatch one shape onto its payload, degrading rather than raising."""
    geometry = _geometry(shape, slide_width, slide_height)
    shape_type = _shape_type(shape)

    # Checked first: a picture must never be previewed as anything else.
    if shape_type in _PICTURE_TYPES:
        return {"kind": "unsupported", **geometry}

    try:
        if shape_type == MSO_SHAPE_TYPE.TABLE or getattr(shape, "has_table", False):
            return _table_payload(shape, geometry)
        if shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
            return _text_payload(shape, geometry)
        if shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            return _shape_payload(shape, geometry)
        # Nothing the builder emits reaches here. Fall back on what the shape
        # can actually do rather than on what it is called.
        if getattr(shape, "has_text_frame", False):
            return _text_payload(shape, geometry)
    except Exception:  # noqa: BLE001 - one odd shape may not lose the others
        return {"kind": "unsupported", **geometry}

    return {"kind": "unsupported", **geometry}


def _background_color(slide: Any) -> str | None:
    """The slide's explicit RGB background, or ``None`` when it is inherited.

    NOTE for callers: this read is not free of side effects. python-pptx
    materializes a ``<p:bg>`` element on a slide that had none, so a deck must
    be SAVED before its preview is derived - see ``converter.preview_for``.
    """
    try:
        return _solid_fill(slide.background.fill)
    except Exception:  # noqa: BLE001
        return None


# --------------------------------------------------------------------------- #
# the payload
# --------------------------------------------------------------------------- #
def build_preview(presentation: Any) -> dict[str, Any]:
    """Describe ``presentation``'s first slide as the pinned preview payload.

    Args:
        presentation: a built :class:`pptx.presentation.Presentation`, already
            saved (see :func:`_background_color`).

    Returns:
        ``{"slide": {...}, "background_color": ... , "shapes": [...]}`` with the
        shapes already in z-order, back to front.
    """
    slide_width = int(getattr(presentation, "slide_width", 0) or 0)
    slide_height = int(getattr(presentation, "slide_height", 0) or 0)

    slides = list(presentation.slides)
    slide = slides[0] if slides else None

    shapes: list[dict[str, Any]] = []
    if slide is not None:
        # Document order IS the z-order, so no sorting happens here - sorting
        # would be the bug.
        for shape in slide.shapes:
            shapes.append(_one_shape(shape, slide_width, slide_height))

    return {
        "slide": {
            "width_in": round(Emu(slide_width).inches, 3),
            "height_in": round(Emu(slide_height).inches, 3),
            "width_pt": round(Emu(slide_width).pt, _PT_DIGITS),
            "height_pt": round(Emu(slide_height).pt, _PT_DIGITS),
        },
        "background_color": _background_color(slide),
        "shapes": shapes,
    }
